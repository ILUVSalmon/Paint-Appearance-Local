"""RO16 -- a Wave Scan sheet with no serial that holds many cars one after
another loads as many cars (every plant).

Claims:

* A blank-serial sheet is split into walks at each row whose checkzone
  resolves to the first point of the plant x model's live Wave Scan list
  (D-RO16-1). With more than one walk, every blank-serial row carries a
  ``car_key`` (D-RO16-2): five walks of one day, model and colour are five
  cars, five ids, 160 rows and no duplicate; the same readings uploaded again
  under another name are all duplicates.
* A walk carrying two colours is held, with a sentence naming it (D-RO16-3).
* A one-walk sheet and a sheet with serials are exactly today's: no
  ``car_key``, the same id (D-RO16-6); a stored AAT / FTM id is unchanged.
* ``v_car_completeness`` counts the five cars, each complete at 32
  (D-RO16-4); the app role inserts ``car_key`` and still cannot UPDATE or
  DELETE ``wavescan``.

Synthetic workbooks only, on FVL's live CX743 list (read, never written).
Every batch a test makes is removed by ``film_cleanup`` (facts first).
"""

from __future__ import annotations

import hashlib
import io
import uuid
from datetime import date

import pytest
from openpyxl import Workbook

from app.domain.cleaner import dedup_key
from app.domain.workbooks import identity_of, set_key, vehicle_id
from tests.test_t04_ingest import admin, engineer, ingest_apps  # noqa: F401
from tests.test_t46_filmthickness import commit, film_cleanup, sent  # noqa: F401

pytestmark = pytest.mark.ingest

PLANT = "FVL"
MODEL = "CX743"
#: FVL's own header row (profile 44): Orientation and Panel2 are ignored.
HEADERS = ["Date", "Year", "Month", "Style", "Color", "Orientation", "Panel",
           "LW", "SW", "du", "Wa", "Wb", "Wc", "Wd", "We", "CF", "Doi", "Panel2"]


@pytest.fixture(scope="module")
def walk(admin_conn) -> list[str]:
    """FVL CX743's live, non-extra list in walk order (32 points)."""
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT zone_code FROM ref_plant_ws_point WHERE plant_code = %s "
            "AND model_code = %s AND retired_at IS NULL AND NOT is_extra "
            "ORDER BY point_no",
            (PLANT, MODEL),
        )
        points = [r[0] for r in cur.fetchall()]
    assert len(points) == 32, points
    return points


@pytest.fixture(scope="module")
def colours(admin_conn) -> list[str]:
    """Two colours FVL offers, as their catalogue names."""
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT color_name FROM ref_plant_color WHERE plant_code = %s "
            "AND is_offered ORDER BY color_name LIMIT 2",
            (PLANT,),
        )
        found = [r[0] for r in cur.fetchall()]
    assert len(found) == 2
    return found


def a_day() -> date:
    """A day no real car was measured on."""
    n = uuid.uuid4().int
    return date(2001 + n % 15, 1 + (n // 15) % 12, 1 + (n // 180) % 28)


def sheet(cars, *, serials=None) -> bytes:
    """One FVL tab, cars one after another. ``cars`` is
    ``[(labels, colour or [colour per label], day)]``. Readings differ per car
    and per test (``base``)."""
    base = uuid.uuid4().int % 1000 / 100
    wb = Workbook()
    ws = wb.active
    ws.title = "CF data template"
    ws.append(HEADERS + (["Serial No"] if serials else []))
    for k, (labels, colour, day) in enumerate(cars):
        names = [colour] * len(labels) if isinstance(colour, str) else colour
        for i, (label, name) in enumerate(zip(labels, names)):
            ws.append(
                [day, day.year, day.month, MODEL, name, "H", label,
                 10.0 + base + k, 20.0 + i / 10, 1.0, 9.0, 17.0, 14.0, 18.0, 9.0,
                 60.0 + i / 10, 80.0, label]
                + ([serials[k]] if serials else [])
            )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def renamed(data: bytes) -> bytes:
    """The same readings in other bytes (a renamed, re-saved copy)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    wb.properties.title = uuid.uuid4().hex
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def upload_one(engineer, film_cleanup, admin_conn, data):  # noqa: F811
    [review] = sent(
        engineer, film_cleanup, admin_conn, f"ro16-{uuid.uuid4().hex[:8]}.xlsx",
        data, PLANT,
    ).values()
    return review


def review_rows(admin_conn, load_id: int) -> list[tuple]:
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT source_row, disposition, reason, detail, canonical->>'car_key' "
            "  FROM stg_review_row WHERE load_id = %s ORDER BY source_row",
            (load_id,),
        )
        return cur.fetchall()


def vehicles(engineer, review) -> list[dict]:  # noqa: F811
    batch = review["batch"]
    body = engineer.get(
        f"/api/v1/plants/{PLANT}/workbooks/{batch['file_hash']}/vehicles",
        params={"load_id": batch["load_id"], "limit": 200},
    ).json()
    return body["items"]


def sha(*parts) -> str:
    return hashlib.sha256("|".join(str(p) for p in parts).encode()).hexdigest()


# =====================================================================
# pure rules
# =====================================================================
def test_the_dedup_key_and_the_id_take_the_car_key_only_when_present():
    day = date(2026, 8, 4)
    # A serial is the car: unchanged.
    assert dedup_key("FVL", "S1", "1HOOD", day, MODEL, "X") == ("FVL", "S1", "1HOOD")
    # Blank serial: the walk joins the key, and two walks are two keys.
    assert dedup_key("FVL", None, "1HOOD", day, MODEL, "X") == ("FVL", "2026-08-04", MODEL, "X", "1HOOD")
    one = dedup_key("FVL", None, "1HOOD", day, MODEL, "X", "a" * 16)
    two = dedup_key("FVL", None, "1HOOD", day, MODEL, "X", "b" * 16)
    assert one != two and len(one) == 6
    # The review id is today's when there is no car_key.
    canonical = {"measured_date": "2026-08-04", "model": MODEL, "color": "X",
                 "serial_no": None}
    key = set_key("LOAD", canonical)
    assert key[4] is None
    assert vehicle_id(7, identity_of(canonical)) == sha(7, "2026-08-04", MODEL, "X", "")
    keyed = set_key("LOAD", {**canonical, "car_key": "a" * 16})
    assert vehicle_id(7, keyed) == sha(7, "2026-08-04", MODEL, "X", "", "a" * 16)


def test_a_known_stored_vehicle_id_is_unchanged(admin_conn):
    """Every stored AAT / FTM History car still has today's id."""
    from app.domain.history import group_vehicles

    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT w.load_id, w.measured_date, w.model, w.color, w.serial_no, "
            "       NULL::text AS car_key, count(*) AS row_count, "
            "       array_agg(DISTINCT w.checkzone) AS zones "
            "  FROM v_wavescan_active w WHERE w.plant IN ('AAT', 'FTM') "
            " GROUP BY 1, 2, 3, 4, 5 ORDER BY 1, 2, 3, 4, 5 LIMIT 50"
        )
        names = [d.name for d in cur.description]
        rows = [dict(zip(names, r)) for r in cur.fetchall()]
    assert rows, "no stored AAT / FTM car to compare"
    for car in group_vehicles(rows, {}):
        assert car.vehicle_id == sha(
            car.load_id, car.measured_date.isoformat(), car.model or "",
            car.color or "", car.serial_no or "",
        )


def test_the_app_role_inserts_the_car_key_and_never_updates_or_deletes(admin_conn):
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT has_column_privilege('app_rw', 'wavescan', 'car_key', 'INSERT'), "
            "       has_column_privilege('app_rw', 'wavescan', 'car_key', 'UPDATE'), "
            "       has_table_privilege('app_rw', 'wavescan', 'UPDATE'), "
            "       has_table_privilege('app_rw', 'wavescan', 'DELETE')"
        )
        assert cur.fetchone() == (True, False, False, False)


# =====================================================================
# the sheet
# =====================================================================
def test_five_cars_on_one_day_of_one_colour_are_five_cars_and_a_copy_is_all_duplicate(
    engineer, admin, admin_conn, film_cleanup, walk, colours  # noqa: F811
):
    day = a_day()
    data = sheet([(walk, colours[0], day)] * 5)
    review = upload_one(engineer, film_cleanup, admin_conn, data)
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert len(rows) == 160
    assert {r[1] for r in rows} == {"LOAD"}, sorted({(r[1], r[3], r[4]) for r in rows}, key=str)[:6]
    keys = [r[4] for r in rows]
    assert len(set(keys)) == 5 and None not in keys
    # Each car's 32 rows carry one key, in file order.
    assert [keys[i * 32] for i in range(5)] == [keys[i * 32 + 31] for i in range(5)]
    found = vehicles(engineer, review)
    assert len(found) == 5 and len({v["vehicle_id"] for v in found}) == 5
    assert all(v["point_count"] == 32 and v["identity_complete"] for v in found)

    commit(engineer, admin, review)
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT count(*), count(DISTINCT car_key), count(*) FILTER "
            "(WHERE serial_no IS NULL) FROM wavescan WHERE load_id = %s",
            (review["batch"]["load_id"],),
        )
        assert cur.fetchone() == (160, 5, 160)
        cur.execute(
            "SELECT zones_recorded, zones_expected, completeness, car_color "
            "  FROM v_car_completeness WHERE plant = %s AND car_key = ANY(%s)",
            (PLANT, list(set(keys))),
        )
        cars = cur.fetchall()
    assert sorted(cars) == [(32, 32, "COMPLETE", colours[0])] * 5

    again = upload_one(engineer, film_cleanup, admin_conn, renamed(data))
    copy = review_rows(admin_conn, again["batch"]["load_id"])
    assert len(copy) == 160 and {r[1] for r in copy} == {"DUPLICATE"}
    assert {r[4] for r in copy} == set(keys)


def test_a_one_walk_sheet_is_exactly_today(
    engineer, admin_conn, film_cleanup, walk, colours  # noqa: F811
):
    day = a_day()
    review = upload_one(
        engineer, film_cleanup, admin_conn, sheet([(walk, colours[0], day)])
    )
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert len(rows) == 32 and {r[1] for r in rows} == {"LOAD"}
    assert {r[4] for r in rows} == {None}
    [car] = vehicles(engineer, review)
    assert car["vehicle_id"] == sha(
        review["batch"]["load_id"], day.isoformat(), MODEL, colours[0], ""
    )


def test_a_sheet_with_serials_is_exactly_today(
    engineer, admin_conn, film_cleanup, walk, colours  # noqa: F811
):
    day = a_day()
    serials = [f"RO16-{uuid.uuid4().hex[:8]}", f"RO16-{uuid.uuid4().hex[:8]}"]
    review = upload_one(
        engineer, film_cleanup, admin_conn,
        sheet([(walk, colours[0], day)] * 2, serials=serials),
    )
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert len(rows) == 64 and {r[1] for r in rows} == {"LOAD"}
    assert {r[4] for r in rows} == {None}
    found = vehicles(engineer, review)
    assert sorted(v["vehicle_id"] for v in found) == sorted(
        sha(review["batch"]["load_id"], day.isoformat(), MODEL, colours[0], s)
        for s in serials
    )


def test_a_walk_carrying_two_colours_is_held_and_the_sentence_names_it(
    engineer, admin_conn, film_cleanup, walk, colours  # noqa: F811
):
    day = a_day()
    mixed = [colours[0]] * 16 + [colours[1]] * 16
    review = upload_one(
        engineer, film_cleanup, admin_conn,
        sheet([(walk, colours[0], day), (walk, mixed, day), (walk, colours[1], day)]),
    )
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert [r[1] for r in rows] == ["LOAD"] * 32 + ["REJECT"] * 32 + ["LOAD"] * 32
    held = {(r[2], r[3]) for r in rows[32:64]}
    [(reason, detail)] = held
    assert reason == "OTHER"
    assert detail.startswith("Car 2 of 3 on this sheet (starting at row 34) carries two colours")
    assert f'"{colours[0]}" and "{colours[1]}"' in detail
    assert "**" not in detail
    # The tab's row says that sentence, not a generic one.
    body = engineer.get(
        f"/api/v1/plants/{PLANT}/workbooks/{review['batch']['file_hash']}/worksheet"
    ).json()
    [tab] = [s for s in body["sheets"] if s["load_id"] == review["batch"]["load_id"]]
    messages = {r["operator_message"] for r in tab["rows"] if r["source_row"] in range(34, 66)}
    assert messages == {detail}
    # The two other cars still load as two cars.
    assert len({r[4] for r in rows if r[1] == "LOAD"}) == 2
