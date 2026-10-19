"""T42 -- parse a WaveScan tab when no template profile matches (L-02, L-03).

Claims, against the running application, worker and database:

* A P703-shaped workbook -- one colour per tab, no DOI / Serial / Paintline
  / B column -- uploaded at AAT is **not** `REJECTED_TEMPLATE` just because
  no profile carries that exact set of optional columns. Its rows are staged
  and reviewable (`PENDING_REVIEW`), and "Yes, load these rows" loads them
  with DOI, serial and paintline NULL. Nothing is invented.
* The worksheet says so in plain words: "This file has no DOI column ...".
* An exact profile is still preferred when one exists (T26 unchanged).
* An extra unknown column still blocks until the Engineer ignores it.
* Two CF columns still block the tab.
* A tab without a required column (Color) is refused and the worksheet
  names that column. A cover tab is still `OUT_OF_SCOPE_SHEET`.
* A Viewer still cannot upload or read the worksheet (403).
* `app_rw` still cannot DELETE (T01).

Every workbook is built in the test. No owner file is read or copied.
"""

from __future__ import annotations

import io
import uuid

import psycopg
import pytest
from openpyxl import Workbook

from app.domain.reader import OPTIONAL_FIELDS, REQUIRED_WITHOUT_PROFILE, describe_sheet
from app.services.worksheet import missing_optional_message
from tests.test_t04_ingest import BASE_ROW, forget_bytes, upload, wait_for
from tests.test_t04_ingest import cleanup  # noqa: F401
from tests.test_t39_uploader_worksheet import (  # noqa: F401  (fixtures)
    HEADERS,
    admin,
    engineer,
    facts,
    ingest_apps,
    viewer,
    ws_url,
)

pytestmark = pytest.mark.ingest

#: The owner's P703-style instrument tab: no DOI, Serial, Paintline, B, and
#: none of the month/year/tolerance helper columns. Chosen so that no seeded
#: or published AAT profile carries exactly this set (asserted per test).
P703_SHAPE = [
    "Date", "Plant", "Model", "Checkzone", "Color",
    "LW", "SW", "CF", "du", "Wa", "Wb", "Wc", "Wd", "We",
]
ZONES = ["01 RRHOOD", "02 RFHOOD", "03 CFHOOD", "04 LFHOOD", "05 LRHOOD"]


# =====================================================================
# helpers
# =====================================================================
def vehicle(count: int, **over) -> list[dict]:
    """One serial-less vehicle on a random past day (ADR 0013 8a dedup key)."""
    n = uuid.uuid4().int
    day = f"{2019 + n % 6}-{(n // 7) % 12 + 1:02d}-{(n // 97) % 28 + 1:02d}"
    return [
        {"Date": day, "Checkzone": ZONES[i % len(ZONES)], "CF": 50.0 + i,
         "Plant": "AAT U704", **over}
        for i in range(count)
    ]


def workbook(tabs: dict[str, tuple[list[str], list[dict]]]) -> bytes:
    """Several tabs, each `(headers, rows)`; values looked up by header."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, (headers, rows) in tabs.items():
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            merged = {**BASE_ROW, **row}
            ws.append([merged.get(h) for h in headers])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sent(engineer, cleanup, admin_conn, name: str, data: bytes) -> dict[str, dict]:
    """Upload at AAT and wait; the review of each tab, keyed by tab name."""
    forget_bytes(admin_conn, data)
    response = upload(engineer, name, data, expected_plant="AAT")
    assert response.status_code == 202, response.text
    out = {}
    for entry in response.json()["batches"]:
        cleanup.append(entry["load_id"])
        out[entry["sheet"]] = wait_for(engineer, entry["load_id"])
    return out


def batch_row(admin_conn, load_id: int) -> dict:
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT status, profile_id, sheet_class, "
            "       (SELECT count(*) FROM stg_raw s WHERE s.load_id = b.load_id) "
            "  FROM load_batch b WHERE load_id = %s",
            (load_id,),
        )
        status, profile_id, sheet_class, staged = cur.fetchone()
    return {"status": status, "profile_id": profile_id,
            "sheet_class": sheet_class, "staged": staged}


def no_profile(admin_conn, review: dict) -> None:
    """The premise of this file: the tab really had no matching profile."""
    got = batch_row(admin_conn, review["batch"]["load_id"])
    assert got["profile_id"] is None, (
        "an AAT profile now carries this exact contract; pick another shape "
        "so this test still exercises the profile-free path"
    )


# =====================================================================
# A. reader contract (no database)
# =====================================================================
def test_the_profile_free_contract_is_required_fields_plus_cf():
    # RO18 (B4): a sheet with no date column is read and asked its month.
    assert REQUIRED_WITHOUT_PROFILE == (
        "plant", "model", "checkzone", "color", "cf",
    )
    for optional in ("doi", "serial_no", "paintline"):
        assert optional in OPTIONAL_FIELDS
        assert optional not in REQUIRED_WITHOUT_PROFILE


def test_missing_optional_is_read_from_the_header_row():
    data = workbook({"Traction Green": (P703_SHAPE, vehicle(1))})
    alias = {h.lower(): h.lower() for h in P703_SHAPE}
    template = describe_sheet(data, "Traction Green", alias)
    assert template.missing_optional[:3] == ("doi", "serial_no", "paintline")
    assert "b" in template.missing_optional
    assert template.missing_for_profile_free == ()
    assert template.as_json()["missing_optional"] == list(template.missing_optional)


def test_the_notice_is_a_plain_sentence():
    assert missing_optional_message([]) is None
    assert missing_optional_message(["doi"]) == (
        "This file has no DOI column. Rows can still load with DOI empty."
    )
    several = missing_optional_message(["doi", "serial_no", "paintline"])
    assert several.startswith("This file has no DOI, Serial No or Paintline column.")
    assert "doi" not in several and "serial_no" not in several, "no raw tokens"


# =====================================================================
# B. the P703 shape at AAT
# =====================================================================
def test_colour_tabs_without_doi_parse_stage_and_load(
    engineer, cleanup, admin_conn
):
    data = workbook({
        "Instructions": (["How to use this workbook"], [{}]),
        "Traction Green": (P703_SHAPE, vehicle(3)),
        "Traction Green 2": (P703_SHAPE, vehicle(2)),
    })
    reviews = sent(engineer, cleanup, admin_conn, "t42-p703-shape.xlsx", data)

    assert "Instructions" not in reviews, "side sheets are workbook evidence, not data batches"

    for tab, expected in (("Traction Green", 3), ("Traction Green 2", 2)):
        review = reviews[tab]
        no_profile(admin_conn, review)
        assert review["batch"]["status"] == "PENDING_REVIEW", review["batch"]["notes"]
        assert review["batch"]["counts"]["rows_loaded"] == expected, review["reject_reasons"]
        assert "UNKNOWN_TEMPLATE" not in review["commit_blockers"]
        assert "UNVERIFIED_PROFILE" not in review["commit_blockers"]
        got = batch_row(admin_conn, review["batch"]["load_id"])
        assert got["staged"] == expected, "rows are staged for review"

    file_hash = reviews["Traction Green"]["batch"]["file_hash"]
    sheet = engineer.get(ws_url("AAT", file_hash)).json()
    assert "doi" in sheet["missing_optional"]
    assert sheet["missing_optional_message"].startswith("This file has no DOI")
    assert "load with" in sheet["missing_optional_message"]
    assert sheet["missing_optional_message"] in sheet["checklist"]
    by_tab = {s["sheet"]: s for s in sheet["sheets"]}
    assert "Instructions" not in by_tab
    file = engineer.get(f"/api/v1/plants/AAT/workbooks/{file_hash}").json()["workbook"]
    assert file["skipped_sheet_count"] == 1
    assert file["skipped_sheets"][0]["source_sheet"] == "Instructions"
    assert file["skipped_sheets"][0]["reason"]
    assert by_tab["Traction Green"]["state"] == "ready"
    assert "doi" in by_tab["Traction Green"]["missing_optional"]
    assert sheet["ready_sheet_count"] == 2

    loaded = engineer.post(ws_url("AAT", file_hash, "load-ready"))
    assert loaded.status_code == 200, loaded.text
    assert loaded.json()["measurements_inserted"] == 5
    stored = facts(admin_conn, reviews["Traction Green"]["batch"]["load_id"])
    assert len(stored) == 3
    assert {(r[4], r[5], r[6]) for r in stored} == {(None, None, None)}, (
        "DOI, serial and paintline are NULL -- nothing invented"
    )
    assert {r[7] for r in stored} == {"AAT"}
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT after ->> 'profile_free' FROM audit_event "
            " WHERE action = 'batch.commit' AND entity_id = %s",
            (str(reviews["Traction Green"]["batch"]["load_id"]),),
        )
        assert cur.fetchone()[0] == "true", "the commit says which contract it used"


def test_a_tab_with_no_plant_column_takes_the_page_plant(
    engineer, cleanup, admin_conn
):
    headers = [h for h in P703_SHAPE if h != "Plant"]
    reviews = sent(engineer, cleanup, admin_conn, "t42-no-plant.xlsx",
                   workbook({"Traction Green": (headers, vehicle(2))}))
    review = reviews["Traction Green"]
    no_profile(admin_conn, review)
    assert review["batch"]["status"] == "PENDING_REVIEW", review["batch"]["notes"]
    assert review["batch"]["counts"]["rows_loaded"] == 2


def test_an_exact_profile_is_still_preferred(engineer, cleanup, admin_conn):
    """The T04 canonical layout has a verified AAT profile (T26 unchanged)."""
    reviews = sent(engineer, cleanup, admin_conn, "t42-exact.xlsx",
                   workbook({"PasteData": (HEADERS, vehicle(1, Plant="AAT"))}))
    review = reviews["PasteData"]
    assert batch_row(admin_conn, review["batch"]["load_id"])["profile_id"] is not None


# =====================================================================
# C. what still blocks
# =====================================================================
def test_an_extra_column_still_needs_ignore(engineer, cleanup, admin_conn):
    extra = f"Operator Note {uuid.uuid4().hex[:6]}"
    rows = vehicle(2, **{extra: "ok"})
    reviews = sent(engineer, cleanup, admin_conn, "t42-extra.xlsx",
                   workbook({"Traction Green": (P703_SHAPE + [extra], rows)}))
    review = reviews["Traction Green"]
    no_profile(admin_conn, review)
    assert review["batch"]["status"] == "PENDING_REVIEW"
    assert "UNREVIEWED_HEADER" in review["commit_blockers"], "never auto-ignored"
    file_hash = review["batch"]["file_hash"]

    sheet = engineer.get(ws_url("AAT", file_hash)).json()
    assert sheet["extra_columns"] == [extra.lower()]
    assert sheet["sheets"][0]["state"] == "needs_answer"
    refused = engineer.post(ws_url("AAT", file_hash, "load-ready"))
    assert refused.status_code == 409, "Yes does not load past an extra column"
    assert facts(admin_conn, review["batch"]["load_id"]) == []

    ignored = engineer.post(ws_url("AAT", file_hash, "ignore-columns"),
                            json={"headers": [extra]})
    assert ignored.status_code == 200, ignored.text
    after = wait_for(engineer, review["batch"]["load_id"])
    assert "UNREVIEWED_HEADER" not in after["commit_blockers"]
    assert after["batch"]["status"] == "PENDING_REVIEW", "ignoring never commits"
    loaded = engineer.post(ws_url("AAT", file_hash, "load-ready"))
    assert loaded.status_code == 200, loaded.text
    assert len(facts(admin_conn, review["batch"]["load_id"])) == 2


def test_two_cf_columns_still_block(engineer, cleanup, admin_conn):
    wb = Workbook()
    ws = wb.active
    ws.title = "Traction Green"
    ws.append(P703_SHAPE + ["CF"])
    for row in vehicle(2):
        merged = {**BASE_ROW, **row}
        ws.append([merged.get(h) for h in P703_SHAPE] + [99.0])
    buffer = io.BytesIO()
    wb.save(buffer)
    reviews = sent(engineer, cleanup, admin_conn, "t42-two-cf.xlsx", buffer.getvalue())
    review = reviews["Traction Green"]
    assert review["batch"]["status"] == "REJECTED_TEMPLATE"
    assert batch_row(admin_conn, review["batch"]["load_id"])["staged"] == 0
    file_hash = review["batch"]["file_hash"]
    sheet = engineer.get(ws_url("AAT", file_hash)).json()
    assert sheet["sheets"][0]["state"] == "needs_answer"
    assert engineer.post(ws_url("AAT", file_hash, "load-ready")).status_code == 409
    assert facts(admin_conn, review["batch"]["load_id"]) == []


def test_a_tab_without_a_colour_column_is_refused_and_says_so(
    engineer, cleanup, admin_conn
):
    headers = [h for h in P703_SHAPE if h != "Color"]
    reviews = sent(engineer, cleanup, admin_conn, "t42-no-colour.xlsx",
                   workbook({"Sheet1": (headers, vehicle(2))}))
    review = reviews["Sheet1"]
    assert review["batch"]["status"] == "REJECTED_TEMPLATE"
    assert batch_row(admin_conn, review["batch"]["load_id"])["staged"] == 0
    sheet = engineer.get(ws_url("AAT", review["batch"]["file_hash"])).json()
    tab = sheet["sheets"][0]
    assert tab["state"] == "blocked"
    assert "no Color column" in tab["operator_message"]
    assert "Admin" not in tab["operator_message"]


# =====================================================================
# D. roles and T01
# =====================================================================
def test_a_viewer_still_cannot_upload_or_read_the_worksheet(
    viewer, engineer, cleanup, admin_conn
):
    data = workbook({"Traction Green": (P703_SHAPE, vehicle(1))})
    forget_bytes(admin_conn, data)
    assert upload(viewer, "t42-viewer.xlsx", data).status_code == 403
    reviews = sent(engineer, cleanup, admin_conn, "t42-viewer.xlsx", data)
    file_hash = reviews["Traction Green"]["batch"]["file_hash"]
    assert viewer.get(ws_url("AAT", file_hash)).status_code == 403
    assert viewer.post(ws_url("AAT", file_hash, "load-ready")).status_code == 403


@pytest.mark.parametrize("table", ["wavescan", "load_batch", "stg_raw"])
def test_app_rw_still_cannot_delete(app_conn, table):
    with app_conn.cursor() as cur:
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            cur.execute(f"DELETE FROM {table}")


def test_openapi_publishes_missing_optional(viewer):
    schemas = viewer.get("/openapi.json").json()["components"]["schemas"]
    assert {"missing_optional", "missing_optional_message"} <= set(
        schemas["WorksheetOut"]["properties"]
    )
    assert "missing_optional" in schemas["WorksheetSheetOut"]["properties"]
