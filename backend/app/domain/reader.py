"""Read a WaveScan workbook.

Ported from ``Paint Appearance/Component/Database/src/wavescan_loader/reader.py``
(`AGENTS.md` rule 4: the logic moves, it is not retyped from memory). Three
things changed, and each is a rule from this repo's specs rather than a
preference:

1. **The alias map is injected.** The vault's module-level ``HEADER_ALIASES``
   dict is now seed data in ``ref_column_alias`` (`01-DATA-MODEL.md` §6.3), so
   a new header spelling is one approved row, not a code change.
   ``AGENTS.md`` §4 lists hardcoding the dict here as a defect.
2. **Bytes, not a path.** Ingest hashes the upload, stores the bytes, and
   parses *those* bytes -- replay reads the same object back
   (`02-PIPELINE.md` §8). Nothing here touches the filesystem.
3. **A template fingerprint** is produced for every sheet (`02-PIPELINE.md`
   §4), so an unknown layout is refused before any data row is parsed.

This module does no cleaning and no validation. It returns whatever the file
said, plus the exact source coordinates, so a bad value can always be traced
back to a cell.

Pure: no FastAPI, no psycopg, no SQLAlchemy. ``openpyxl`` is a file parser,
not a framework, and the dependency rule in `00-ARCHITECTURE.md` §5 names the
three that matter.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import IO, Collection, Iterable, Iterator, Mapping, Sequence

from openpyxl import load_workbook

from app.domain import ods
from app.domain.film_header import (
    DOMAIN_FILMTHICKNESS,
    EC1_LOCATION_HEADER,
    HEADER_SEARCH_ROWS,
    IGNORED_FIELD,
    SHEET_FILMTHICKNESS,
    classify_film,
    film_missing,
    layer_block_base_header,
    layer_block_before,
    parse_ec1_banner,
    parse_pelt_meta,
)

#: Logical fields a sheet must map before any data row is parsed.
#: (`02-PIPELINE.md` §4: "Required field missing -> REJECTED_TEMPLATE".)
#: ADR 0013: ``serial_no`` is not here. A sheet with no Serial column is a
#: sheet whose serials are blank, and blank serial loads NULL.
#: RO18 (B4, RO09): ``date`` is not required. A sheet with no date column is
#: read, and every row waits on one question for the sheet -- which month?
REQUIRED_FIELDS = ("plant", "model", "checkzone", "color")

#: The ten instrument measures, in ``wavescan`` column order. Dedup compares
#: over these and the measurement reads zip them; neither is a load rule.
MEASURES_CORE = ("lw", "sw", "cf", "doi", "du", "wa", "wb", "wc", "wd", "we")
#: ADR 0013: what a point needs to load. CF always; LW and SW unless an
#: Engineer approved that point's holes (``MEASURES_HOLE_APPROVABLE``). DOI,
#: DU and Wa-We are optional and load NULL when blank or absent.
MEASURES_REQUIRED = ("lw", "sw", "cf")
MEASURES_HOLE_APPROVABLE = ("lw", "sw")
MEASURE_ALL = MEASURES_CORE + ("b",)

#: T98 (D-98-1): the one sentence ``Cleaner`` writes when a point is short of
#: a required measure, and the one reader of it -- the review screen must name
#: what is actually missing, and nothing loads without CF.
MISSING_MEASUREMENT_DETAIL = "required measurement(s) missing: "


def missing_measures(detail: str | None) -> tuple[str, ...]:
    """The required measures a ``MISSING_MEASUREMENT`` detail names."""
    _, found, names = (detail or "").partition(MISSING_MEASUREMENT_DETAIL)
    if not found:
        return ()
    named = {n.strip().lower() for n in names.split(",")}
    return tuple(m for m in MEASURES_REQUIRED if m in named)


#: T42 (L-02): what a WaveScan tab must map to be read when **no** template
#: profile claims its exact contract. The profile-free contract is ADR 0013
#: decision 1 at the column level: the required identity fields plus CF.
#: `plant` may be supplied by the upload page. LW/SW are not here: a sheet
#: without them parses and its points hold for hole approval, row by row.
REQUIRED_WITHOUT_PROFILE = REQUIRED_FIELDS + ("cf",)
#: T42: fields a sheet may simply not have. Their absence is a notice on the
#: worksheet (the values load NULL), never a new template. Order is the
#: order they are named to a person.
OPTIONAL_FIELDS = (
    "doi", "serial_no", "paintline", "du", "wa", "wb", "wc", "wd", "we", "b",
)

#: The WaveScan vocabulary (T26 / spec S1, S6). T46 adds FilmThickness as a
#: second **package** (`film`, `film_header`); this file only learns where a
#: FilmBuild header row sits and which domain a template was read under, so
#: an alias, a profile and a batch all say which vocabulary they used.
DOMAIN_WAVESCAN = "WAVESCAN"

#: What a sheet's header row turned out to be (T26 / spec S10). Routing is by
#: **headers**, never by the tab's name: a workbook that calls its data sheet
#: "Sheet1" and its instructions "CF Data" would otherwise be read backwards.
SHEET_WAVESCAN = "WAVESCAN"
SHEET_OUT_OF_SCOPE = "OUT_OF_SCOPE_SHEET"
SHEET_UNKNOWN_DATA = "UNKNOWN_DATA_SHEET"

#: A WaveScan header row carries measures. Three is the smallest number that
#: cannot be reached by a colour or film-build sheet that happens to share a
#: letter with one of ours, and every real CF layout reviewed carries eleven.
WAVESCAN_MIN_MEASURES = 3
#: ...and it says which vehicle, panel or day the measures belong to. The FVL
#: 18-column contract has no Plant, Serial or Paintline, so none of those
#: three may be required to recognise the *kind* of sheet.
WAVESCAN_IDENTITY_FIELDS = ("checkzone", "model", "serial_no", "date")

#: Every logical field the staging table has a column for.
STAGED_FIELDS = (
    "date", "month", "year", "plant", "model", "checkzone", "color",
    "tolerance_group", "paintline", "serial_no",
) + MEASURE_ALL


class ReaderError(RuntimeError):
    """The workbook cannot be read at all."""


class TemplateError(ReaderError):
    """The sheet's headers do not carry the required fields.

    Carries the diff a reviewer needs: what was recognised, what is missing,
    and the header row verbatim.
    """

    def __init__(self, message: str, template: "SheetTemplate") -> None:
        super().__init__(message)
        self.template = template


class DuplicateFieldError(TemplateError):
    """Two columns of one sheet mean the same logical field (T26 / spec S8).

    A subclass of `TemplateError` because the answer is the same -- refuse
    before a data row is interpreted -- but the *reason* is the opposite of a
    missing column: the sheet says a thing twice and nothing in this platform
    may pick which one it meant. Until T26 the later column was silently
    dropped, which is how a swapped measure column could have become a wrong
    measurement nobody could see.
    """


def sha256_hex(data: bytes) -> str:
    """Content identity. `02-PIPELINE.md` §3: a file *is* its bytes."""
    return hashlib.sha256(data).hexdigest()


def normalise_header(cell_value: object) -> str:
    """lower / trim / underscore -> space, exactly as `01-DATA-MODEL.md` §6.3
    describes ``ref_column_alias.raw_header``."""
    return str(cell_value or "").strip().lower().replace("_", " ")


@dataclass(frozen=True)
class SheetTemplate:
    """What a sheet's header row says, and what it was understood to mean."""

    sheet: str
    raw_headers: tuple[str, ...]
    recognised: dict[int, str]
    unrecognised: tuple[str, ...]
    fingerprint: str
    #: Required fields no header carries but an approved template profile
    #: supplies (T18). Today that is `plant`, filled from the upload's
    #: `expected_plant` when the profile says `fill_missing_plant`. It is a
    #: decision made *about* this layout, not a reading of it: it changes
    #: whether the sheet is parseable and deliberately **not** the
    #: fingerprint, so the same headers keep identifying the same template
    #: whether or not a profile exists for them.
    supplied: tuple[str, ...] = ()
    #: Which vocabulary the headers above were read under (T26 / spec S6).
    #: `WAVESCAN` throughout this slice; carried so a future FilmThickness
    #: package cannot inherit this one's meaning of `Panel` or `Date`.
    domain: str = DOMAIN_WAVESCAN
    #: `(column index, logical field)` for every column **after** the first
    #: that means an already-recognised field. Never silently dropped: a
    #: sheet that says `CF` twice is ambiguous and is refused (spec S8).
    duplicates: tuple[tuple[int, str], ...] = ()
    #: Headers an approved profile column policy marks IGNORED, as actually
    #: seen on this sheet. They are neither recognised nor unreviewed: a
    #: named owner has already said this column is not a WaveScan field
    #: (spec 5.3 -- this is how FVL `Panel2` becomes auxiliary).
    ignored: tuple[str, ...] = ()
    #: T46. 1-based spreadsheet row the header was found on. WaveScan is
    #: always row 1; a PELT FilmBuild report puts it under its preamble.
    header_row: int = 1
    #: T46. PELT ``Label:`` -> value pairs above the header (``Date``,
    #: ``Job Number``, ``Vin``, ``Layer 1``..``Layer 6``). Empty for WaveScan.
    meta: tuple[tuple[str, object], ...] = ()
    #: T46. The Base column was bound because PELT ``Layer 2`` names it.
    base_from_layer_block: bool = False
    #: RO02. The Primer column was bound because it is the PELT layer
    #: directly before E-coat (D-RO02-1).
    primer_from_layer_block: bool = False

    @property
    def meta_map(self) -> dict[str, object]:
        return dict(self.meta)

    @property
    def missing_required(self) -> tuple[str, ...]:
        if self.domain == DOMAIN_FILMTHICKNESS:
            return film_missing(self.recognised.values(), self.meta_map)
        present = set(self.recognised.values()) | set(self.supplied)
        return tuple(f for f in REQUIRED_FIELDS if f not in present)

    @property
    def missing_for_profile_free(self) -> tuple[str, ...]:
        """Required columns a sheet read without a profile still lacks (T42)."""
        if self.domain == DOMAIN_FILMTHICKNESS:
            return self.missing_required
        present = set(self.recognised.values()) | set(self.supplied)
        return tuple(f for f in REQUIRED_WITHOUT_PROFILE if f not in present)

    @property
    def missing_optional(self) -> tuple[str, ...]:
        """Optional fields no column carries (T42). A notice, not a refusal."""
        if self.domain == DOMAIN_FILMTHICKNESS:
            return ()
        present = set(self.recognised.values())
        return tuple(f for f in OPTIONAL_FIELDS if f not in present)

    @property
    def duplicate_fields(self) -> tuple[str, ...]:
        """Logical fields more than one column claims, sorted."""
        return tuple(sorted({field_ for _, field_ in self.duplicates}))

    @property
    def semantic_signature(self) -> str:
        """The order-independent contract this sheet carries (spec S7).

        Two workbooks with the same logical fields in different column order
        share this hash and differ on `fingerprint`. That is the whole of
        T26's selection change: a reorder is a change of *layout*, not of
        *contract*, and refusing it was the T18 brittleness this slice fixes.
        """
        return semantic_signature(self.recognised.values(), self.domain)

    @property
    def unreviewed_headers(self) -> tuple[str, ...]:
        """Extras nobody has mapped and nobody has approved ignoring.

        Not the same as `unrecognised`, which is already filtered: a header a
        named owner marked IGNORED on this profile is a decision, and a
        header with no decision at all is `UNREVIEWED_HEADER` (spec S9).
        """
        return self.unrecognised

    @property
    def sheet_class(self) -> str:
        """`WAVESCAN`, `FILMTHICKNESS` (T46), `OUT_OF_SCOPE_SHEET` or
        `UNKNOWN_DATA_SHEET` (S10)."""
        if self.domain == DOMAIN_FILMTHICKNESS:
            if classify_film(self.recognised.values()):
                return SHEET_FILMTHICKNESS
            return SHEET_UNKNOWN_DATA if self.recognised else SHEET_OUT_OF_SCOPE
        return classify_sheet(self.recognised.values())

    @property
    def is_parseable(self) -> bool:
        """Can a data row of this sheet be interpreted at all?

        False when a required field is missing **or** when two columns claim
        one field. The two are different refusals with the same consequence:
        no data row is read, because nothing here may choose a meaning.
        """
        return not self.missing_required and not self.duplicates

    def as_json(self) -> dict:
        """Stored verbatim on ``load_batch.raw_headers`` (`02-PIPELINE.md` §4)."""
        return {
            "sheet": self.sheet,
            "domain": self.domain,
            "headers": list(self.raw_headers),
            "recognised": {str(k): v for k, v in sorted(self.recognised.items())},
            "unrecognised": list(self.unrecognised),
            "fingerprint": self.fingerprint,
            "semantic_signature": self.semantic_signature,
            "sheet_class": self.sheet_class,
            "missing_required": list(self.missing_required),
            "missing_optional": list(self.missing_optional),
            "duplicate_fields": list(self.duplicate_fields),
            "duplicate_columns": [
                {"column": idx, "logical_field": field_}
                for idx, field_ in self.duplicates
            ],
            "ignored_headers": list(self.ignored),
            "supplied_by_profile": list(self.supplied),
            "header_row": self.header_row,
            "meta": {
                # T75: a PELT `Time:` cell is a `datetime.time` (.xlsx and .ods).
                k: (v.isoformat() if isinstance(v, (datetime, date, time)) else v)
                for k, v in self.meta
            },
            "base_from_layer_block": self.base_from_layer_block,
            "primer_from_layer_block": self.primer_from_layer_block,
        }


@dataclass
class RawRow:
    """One spreadsheet row, verbatim, with its source coordinates."""

    source_file: str
    source_sheet: str
    source_row: int
    values: dict[str, object]
    #: Cells whose header mapped to no logical field. Kept rather than
    #: dropped: `01-DATA-MODEL.md` §4 puts them in ``stg_raw.extra`` so
    #: nothing a file said is lost just because no alias existed yet.
    extra: dict[str, object] = field(default_factory=dict)

    def get(self, field_name: str) -> object:
        return self.values.get(field_name)

    def text(self, field_name: str) -> str | None:
        """Value as trimmed text, or None if blank."""
        v = self.values.get(field_name)
        if v is None:
            return None
        if isinstance(v, str):
            v = v.strip()
            return v or None
        if isinstance(v, (datetime, date)):
            return v.isoformat()
        return str(v)

    def is_blank(self) -> bool:
        cells = list(self.values.values()) + list(self.extra.values())
        return all(v is None or (isinstance(v, str) and not v.strip()) for v in cells)


def _map_headers(
    header_cells: tuple,
    alias_map: Mapping[str, str],
    ignored_headers: Collection[str] = (),
) -> tuple[dict[int, str], list[str], list[tuple[int, str]], list[str]]:
    """Column index -> logical field, the extras, the repeats, the ignored.

    First match wins, the same rule the vault reader used, so which column a
    field binds to is unchanged. What changed in T26 is that the *later*
    column is no longer thrown away in silence: it is returned as a duplicate
    and the caller refuses the sheet (spec S8). Before T26 a workbook with
    two `CF` columns parsed happily using the first one, which is exactly the
    shape a silently swapped measure would have.

    ``ignored_headers`` are normalised headers an approved profile column
    policy marks IGNORED. They are skipped **before** the alias lookup, so an
    ignored column can neither become a field nor collide with one.
    """
    recognised: dict[int, str] = {}
    unrecognised: list[str] = []
    duplicates: list[tuple[int, str]] = []
    ignored: list[str] = []
    ignore = set(ignored_headers)
    for idx, cell in enumerate(header_cells):
        key = normalise_header(cell)
        if not key:
            continue
        if key in ignore:
            ignored.append(key)
            continue
        logical = alias_map.get(key)
        if logical is None:
            unrecognised.append(key)
        elif logical in recognised.values():
            duplicates.append((idx, logical))
        else:
            recognised[idx] = logical
    return recognised, unrecognised, duplicates, ignored


def _fingerprint(recognised: dict[int, str], unrecognised: list[str]) -> str:
    """Stable hash of the ordered header interpretation.

    Two files with the same columns in the same order fingerprint the same,
    whichever spelling each used -- which is the point: the fingerprint
    identifies a *template*, and an alias is what makes two spellings one
    template.
    """
    ordered = [recognised.get(i, "?") for i in range(max(recognised, default=-1) + 1)]
    payload = "|".join(ordered) + "||" + "|".join(sorted(unrecognised))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def semantic_signature(
    fields: Iterable[str], domain: str = DOMAIN_WAVESCAN
) -> str:
    """Stable hash of the **set** of logical fields a sheet carries (S7).

    Sorted and de-duplicated, so column order cannot reach the hash: FTM
    August and SAP July reorder the same twenty-one fields and signature
    identically, while `_fingerprint` above records that the layout moved.

    The domain is part of the payload rather than a separate column check, so
    a future FilmThickness sheet whose fields happened to be spelled the same
    could never select a WaveScan profile by accident.

    Sorting is plain codepoint order. `ref_template_profile` is backfilled by
    revision 0009 with the same ordering (`ORDER BY ... COLLATE "C"`), so the
    database and this function cannot drift over an underscore.
    """
    payload = domain + "||" + "|".join(sorted(set(fields)))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def classify_sheet(fields: Iterable[str]) -> str:
    """Which kind of sheet this header row belongs to (spec S10).

    Three answers, and the difference between the last two is what a person
    is asked to do about it:

    * **`WAVESCAN`** -- it carries WaveScan measures and says what they
      belong to. Whether it then *parses* is a separate question the profile
      answers; a sheet can be WaveScan and still be `REJECTED_TEMPLATE`.
    * **`OUT_OF_SCOPE_SHEET`** -- its header row recognises **no** field of
      this domain at all: a Methods or instructions sheet, or a grid in a
      vocabulary this domain shares nothing with (Color Data). Preserved as
      evidence. It is **not** a broken WaveScan template and nobody is asked
      to map its headers.
    * **`UNKNOWN_DATA_SHEET`** -- it speaks *some* WaveScan but not enough to
      be one, however few columns it has. That is the ambiguous case, so it
      blocks and asks for a classification. (A column count is deliberately
      not a reason to call a sheet out of scope: a two-column sheet carrying
      `CF` is WaveScan-shaped evidence, and preserving it quietly would be a
      silent drop.)

    The tab's name is deliberately not an input. `02-PIPELINE.md` §4 already
    refuses to let a filename be a key, and a sheet name is the same kind of
    label.
    """
    present = set(fields)
    measures = len(present & set(MEASURE_ALL))
    identifies = bool(present & set(WAVESCAN_IDENTITY_FIELDS))
    if measures >= WAVESCAN_MIN_MEASURES and identifies:
        return SHEET_WAVESCAN
    if not present:
        return SHEET_OUT_OF_SCOPE
    return SHEET_UNKNOWN_DATA


def _open(source: bytes | IO[bytes]):
    from io import BytesIO

    handle = BytesIO(source) if isinstance(source, (bytes, bytearray)) else source
    return load_workbook(handle, read_only=True, data_only=True)


def _as_bytes(source: bytes | IO[bytes]) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    position = source.tell()
    data = source.read()
    source.seek(position)
    return data


def iter_rows(
    source: bytes | IO[bytes], sheet: str, max_row: int | None = None
) -> Iterator[tuple]:
    """Every row of one sheet as a tuple of values, `.xlsx` or `.ods` (T46).

    Row *n* of the iterator is spreadsheet row *n* whichever format the
    bytes are, so a source row number always names the cell a person opens.
    """
    data = _as_bytes(source)
    if ods.is_ods(data):
        yield from ods.iter_rows(data, sheet, max_row)
        return
    wb = _open(data)
    try:
        yield from wb[sheet].iter_rows(values_only=True, max_row=max_row)
    finally:
        wb.close()


def sheet_names(source: bytes | IO[bytes]) -> list[str]:
    data = _as_bytes(source)
    if ods.is_ods(data):
        return ods.sheet_names(data)
    wb = _open(data)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def describe_sheet(
    source: bytes | IO[bytes],
    sheet: str,
    alias_map: Mapping[str, str],
    supplied_fields: Sequence[str] = (),
    ignored_headers: Collection[str] = (),
    domain: str = DOMAIN_WAVESCAN,
) -> SheetTemplate:
    """Read only the header row. No data row is touched.

    This is what `02-PIPELINE.md` §4 requires before deciding whether to
    parse: a sheet whose template is unknown must be refused *without*
    interpreting its rows.

    ``supplied_fields`` are required fields an approved template profile
    provides even though no column carries them (T18). They never affect the
    fingerprint -- a layout is its headers -- only whether this sheet can be
    parsed.

    ``ignored_headers`` come from an approved profile's column policy (T26).
    **They do not move the fingerprint.** The exact ordered-layout hash is
    computed from the header row and the alias map alone, so it stays what it
    has always been -- pure provenance about the bytes -- and marking a
    header ignored later cannot silently rewrite the layout identity of a
    workbook already on file.
    """
    rows = iter_rows(source, sheet, max_row=1)
    try:
        try:
            header = next(rows)
        except StopIteration:
            raise ReaderError(f"sheet [{sheet}] is empty") from None
        # The exact ordered-layout hash: no policy applied, first match wins,
        # later repeats dropped -- byte-identical to T18's algorithm, so the
        # profiles revision 0006 seeded still name the layouts they named.
        exact_recognised, exact_unrecognised, _, _ = _map_headers(header, alias_map)
        recognised, unrecognised, duplicates, ignored = _map_headers(
            header, alias_map, ignored_headers
        )
        return SheetTemplate(
            sheet=sheet,
            raw_headers=tuple(str(h) for h in header if h is not None),
            recognised=recognised,
            unrecognised=tuple(unrecognised),
            fingerprint=_fingerprint(exact_recognised, exact_unrecognised),
            supplied=tuple(supplied_fields),
            domain=domain,
            duplicates=tuple(duplicates),
            ignored=tuple(ignored),
        )
    finally:
        rows.close()


def read_sheet(
    source: bytes | IO[bytes],
    sheet: str,
    alias_map: Mapping[str, str],
    source_file: str,
    supplied_fields: Sequence[str] = (),
    ignored_headers: Collection[str] = (),
    domain: str = DOMAIN_WAVESCAN,
) -> Iterator[RawRow]:
    """Yield every non-blank data row from one sheet.

    Raises ``TemplateError`` if a required field is missing, and
    ``DuplicateFieldError`` if two columns mean one field -- both **before**
    yielding anything, because a sheet whose meaning is ambiguous must not
    have its rows interpreted at all (spec S8). A field an approved profile
    supplies (``supplied_fields``, T18) is not missing: the row simply has no
    value for it, and the cleaner fills it from the upload's
    ``expected_plant``.
    """
    template = describe_sheet(
        source, sheet, alias_map, supplied_fields, ignored_headers, domain
    )
    if template.duplicates:
        raise DuplicateFieldError(
            f"{source_file} [{sheet}] gives more than one column the same "
            f"meaning: "
            + "; ".join(
                f"column {idx + 1} repeats {field_!r}"
                for idx, field_ in template.duplicates
            )
            + ". No data row was read. Nothing in this platform chooses "
            f"which column a repeated field meant -- remove the repeat, or "
            f"have the owner publish a profile whose column policy ignores "
            f"one of them. Header row was: {list(template.raw_headers)}",
            template,
        )
    if not template.is_parseable:
        raise TemplateError(
            f"{source_file} [{sheet}] is missing required column(s): "
            f"{', '.join(template.missing_required)}. "
            f"Recognised: {', '.join(sorted(set(template.recognised.values()))) or '(none)'}. "
            f"Header row was: {list(template.raw_headers)}",
            template,
        )

    yield from _data_rows(source, sheet, template, source_file)


def _data_rows(
    source: bytes | IO[bytes],
    sheet: str,
    template: SheetTemplate,
    source_file: str,
) -> Iterator[RawRow]:
    """Every non-blank row under the header, keyed by logical field."""
    rows = iter_rows(source, sheet)
    try:
        for _ in range(template.header_row):
            next(rows, None)  # preamble and header, already interpreted
        for offset, row in enumerate(rows, start=template.header_row + 1):
            values = {
                logical: row[idx] if idx < len(row) else None
                for idx, logical in template.recognised.items()
            }
            extra = {}
            for idx, cell in enumerate(row):
                if idx in template.recognised or cell is None:
                    continue
                key = normalise_header(
                    template.raw_headers[idx]
                    if idx < len(template.raw_headers)
                    else f"column {idx + 1}"
                )
                extra[key or f"column {idx + 1}"] = cell
            raw = RawRow(
                source_file=source_file,
                source_sheet=sheet,
                source_row=offset,
                values=values,
                extra=extra,
            )
            if raw.is_blank():
                continue
            yield raw
    finally:
        rows.close()


# =====================================================================
# T46: FilmThickness header rows
# =====================================================================
def describe_film_sheet(
    source: bytes | IO[bytes],
    sheet: str,
    alias_map: Mapping[str, str],
    ignored_headers: Collection[str] = (),
) -> SheetTemplate | None:
    """Find a FilmBuild header in the first rows; None when there is none.

    ``alias_map`` is the **FILMTHICKNESS** vocabulary only (spec F3); the
    WaveScan map never reaches this function, and this one never reaches
    `describe_sheet`, so `CC1 mic` cannot mean a WaveScan field and `Panel`
    cannot mean a film location.

    The first row (of `HEADER_SEARCH_ROWS`) that names a location and two
    paint layers is the header. The rows above it are read as the PELT
    preamble. When ``Layer 2:`` names a colour that is also a header, that
    column is Base: the file's own layer block says so (spec 5.5). No data
    row is touched.
    """
    head = list(iter_rows(source, sheet, max_row=HEADER_SEARCH_ROWS))
    for index, header in enumerate(head):
        if not header or not any(isinstance(c, str) and c.strip() for c in header):
            continue
        meta = parse_pelt_meta(head[:index])
        exact = _map_film_headers(header, alias_map, (), meta)
        if not classify_film(exact[0].values()):
            continue
        # T56 (D-56-01): an AAT EC1 / Byk sheet names its location column
        # `Checkzone` and keeps date / colour / sample id in the banner above
        # it, not in `Label:` cells. A PELT label always wins where present.
        location_header = next(
            (normalise_header(header[i]) for i, f in exact[0].items() if f == "location"),
            None,
        )
        if location_header == EC1_LOCATION_HEADER:
            for key, value in parse_ec1_banner(head[:index]).items():
                meta.setdefault(key, value)
        recognised, unrecognised, duplicates, ignored, bound, primer = _map_film_headers(
            header, alias_map, ignored_headers, meta
        )
        return SheetTemplate(
            sheet=sheet,
            raw_headers=tuple(str(h) for h in header if h is not None),
            recognised=recognised,
            unrecognised=tuple(unrecognised),
            fingerprint=_fingerprint(exact[0], exact[1]),
            domain=DOMAIN_FILMTHICKNESS,
            duplicates=tuple(duplicates),
            ignored=tuple(ignored),
            header_row=index + 1,
            meta=tuple(meta.items()),
            base_from_layer_block=bound,
            primer_from_layer_block=primer,
        )
    return None


def _map_film_headers(header, alias_map, ignore, meta):
    """`_map_headers`, plus the Base and Primer columns the PELT layer block names.

    RO02 (D-RO02-1/2): with no alias-recognised primer, the column directly
    left of the E-coat column is Primer when its header is the block's layer
    directly before E-coat. Base binds first, so a pre-E-coat layer named
    like Layer 2 leaves Base on the first such column and Primer on this one.
    Each binding takes exactly one entry out of ``unrecognised`` (D-RO02-3).

    T60 (D-60-05): a second (third ...) column mapped to the non-fact
    ``ignored`` field is not a duplicate -- ``Ground-coat`` and
    ``Orientation`` on one sheet are two named non-facts, not one field said
    twice. The first is kept under ``ignored``; the others stay in the row's
    ``extra`` as evidence. None is ever a paint layer.
    """
    recognised, unrecognised, duplicates, ignored = _map_headers(
        header, alias_map, ignore
    )
    duplicates = [(idx, f) for idx, f in duplicates if f != IGNORED_FIELD]
    base_header = layer_block_base_header(meta)
    bound = False
    if base_header and "base" not in recognised.values():
        for idx, cell in enumerate(header):
            if normalise_header(cell) == base_header and idx not in recognised:
                recognised[idx] = "base"
                if base_header in unrecognised:
                    unrecognised.remove(base_header)
                bound = True
                break
    primer = False
    ecoat = next((i for i, f in recognised.items() if f == "ecoat"), 0)
    if ecoat and "primer" not in recognised.values() and ecoat - 1 not in recognised:
        key = normalise_header(header[ecoat - 1])
        if key in unrecognised and key == layer_block_before(
            meta, normalise_header(header[ecoat])
        ):
            recognised[ecoat - 1] = "primer"
            unrecognised.remove(key)
            primer = True
    return recognised, unrecognised, duplicates, ignored, bound, primer


def read_film_sheet(
    source: bytes | IO[bytes],
    sheet: str,
    template: SheetTemplate,
    source_file: str,
) -> Iterator[RawRow]:
    """Every non-blank row under a FilmBuild header, keyed by logical field.

    Refuses before yielding anything, exactly as `read_sheet` does, when the
    template repeats a field or lacks one (spec S8).
    """
    if template.duplicates:
        raise DuplicateFieldError(
            f"{source_file} [{sheet}] gives more than one column the same "
            f"meaning: "
            + "; ".join(
                f"column {idx + 1} repeats {field_!r}"
                for idx, field_ in template.duplicates
            )
            + ". No data row was read.",
            template,
        )
    if template.missing_required:
        raise TemplateError(
            f"{source_file} [{sheet}] cannot be read as FilmThickness: "
            f"missing {', '.join(template.missing_required)}. "
            f"Header row {template.header_row} was: {list(template.raw_headers)}",
            template,
        )
    yield from _data_rows(source, sheet, template, source_file)
