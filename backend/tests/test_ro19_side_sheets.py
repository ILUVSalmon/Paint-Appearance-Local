"""RO19: side sheets belong to a workbook, never a rejected data batch.

Synthetic plant and uploads exist only in a rolled-back transaction. No
worker, live HTTP upload, suite process or owner bytes are involved.
"""
from __future__ import annotations

import io

import psycopg
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text

from app.core.config import Settings
from app.repositories import batches
from app.services import workbooks
from app.services.ingest import DuplicateUpload, IngestService
from app.services.raw_store import RawStore

PLANT = "ZRO19SIDE"


@pytest.fixture()
def context(admin_dsn, tmp_path):
    engine = create_engine("postgresql+psycopg://", creator=lambda: psycopg.connect(admin_dsn))
    try:
        with engine.connect() as conn:
            tx = conn.begin()
            try:
                conn.execute(text("INSERT INTO ref_plant (plant_code,plant_name,is_active) "
                                  "VALUES (:p,'RO19 side-sheet test',true)"), {"p": PLANT})
                actor = str(conn.execute(text("SELECT id FROM app_user ORDER BY id LIMIT 1")).scalar_one())
                conn.execute(text("SET LOCAL ROLE app_rw"))
                yield conn, IngestService(Settings(), RawStore(tmp_path)), actor
            finally:
                tx.rollback()
    finally:
        engine.dispose()


def workbook(*, data_sheet=False, ambiguous=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Instrument instructions"
    ws.append(["How to use this workbook"])
    ws.append(["Keep this sheet as evidence"])
    wb.create_sheet("Layer diagram").append(["Diagram"])
    if data_sheet:
        ws = wb.create_sheet("Measurements")
        ws.append(["Date", "Plant", "Model", "Checkzone", "Color", "CF", "LW", "SW"])
        ws.append(["2020-01-01", PLANT, "Synthetic model", "Synthetic point", "Synthetic colour", 40, 2, 3])
    if ambiguous:
        ws = wb.create_sheet("Unreadable data")
        ws.append(["CF"])
        ws.append([40])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def accepted(context, data):
    conn, service, actor = context
    return service.accept_upload(conn, [("ro19-side-sheets.xlsx", data)],
                                 expected_plant=PLANT, actor_user_id=actor)


def test_side_sheets_create_one_workbook_record_and_only_a_real_data_batch(context):
    conn, _, _ = context
    outcome = accepted(context, workbook(data_sheet=True))
    assert len(outcome.load_ids) == 1
    stored = conn.execute(text("SELECT skipped_sheets FROM workbook_upload WHERE plant_code=:p"),
                          {"p": PLANT}).scalar_one()
    assert {s["source_sheet"] for s in stored} == {"Instrument instructions", "Layer diagram"}
    assert all(s["reason"] for s in stored)
    assert conn.execute(text("SELECT count(*) FROM load_batch WHERE plant_code=:p"),
                        {"p": PLANT}).scalar_one() == 1
    digest = outcome.file_hash_by_name["ro19-side-sheets.xlsx"]
    item = workbooks.get_workbook(conn, PLANT, digest)
    assert item["sheet_count"] == 1
    assert item["skipped_sheet_count"] == 2
    assert {s["source_sheet"] for s in item["skipped_sheets"]} == {"Instrument instructions", "Layer diagram"}
    assert conn.execute(text("SELECT count(*) FROM audit_event WHERE action='workbook.upload' "
                             "AND entity_type='workbook_upload' AND entity_id=:key"),
                        {"key": f"{PLANT}:{digest}"}).scalar_one() == 1


def test_an_all_skipped_workbook_is_visible_and_cannot_be_uploaded_twice(context):
    conn, _, _ = context
    data = workbook()
    outcome = accepted(context, data)
    assert outcome.load_ids == []
    digest = outcome.file_hash_by_name["ro19-side-sheets.xlsx"]
    items, total = workbooks.list_workbooks(conn, PLANT, limit=20, offset=0)
    assert total == 1
    assert items[0]["file_hash"] == digest
    assert items[0]["operational_summary"] == "SKIPPED"
    assert items[0]["sheet_count"] == 0
    assert items[0]["skipped_sheet_count"] == 2
    history = workbooks.history(conn, PLANT, digest)
    assert any(event["action"] == "workbook.upload" for event in history)
    with pytest.raises(DuplicateUpload):
        accepted(context, data)
    assert conn.execute(text("SELECT count(*) FROM workbook_upload WHERE plant_code=:p"),
                        {"p": PLANT}).scalar_one() == 1


def test_purged_bytes_can_be_uploaded_again_though_the_workbook_row_survives(context):
    """A purge and a plant answer that frees the hash both delete the batches
    and write `workbook.purge`, and neither can remove the `workbook_upload`
    row. Without reading that trace the re-upload is refused 409 for good."""
    conn, _, actor = context
    data = workbook()
    digest = accepted(context, data).file_hash_by_name["ro19-side-sheets.xlsx"]
    with pytest.raises(DuplicateUpload):
        accepted(context, data)

    conn.execute(text("INSERT INTO audit_event(actor_type,action,entity_type,entity_id,after) "
                      "VALUES('user','workbook.purge','workbook',:key,'{}'::jsonb)"),
                 {"key": f"{PLANT}:{digest}"})
    again = accepted(context, data)
    assert again.file_hash_by_name["ro19-side-sheets.xlsx"] == digest
    # The surviving row is kept, not duplicated: the insert cannot update it.
    assert conn.execute(text("SELECT count(*) FROM workbook_upload WHERE plant_code=:p"),
                        {"p": PLANT}).scalar_one() == 1


def test_a_real_unreadable_data_sheet_still_has_a_rejected_batch(context):
    conn, _, _ = context
    outcome = accepted(context, workbook(ambiguous=True))
    assert len(outcome.load_ids) == 1
    batch = batches.get_batch(conn, outcome.load_ids[0])
    assert batch["source_sheet"] == "Unreadable data"
    assert batch["sheet_class"] == "UNKNOWN_DATA_SHEET"
    assert batch["status"] == "REJECTED_TEMPLATE"


def test_legacy_side_batch_is_kept_but_not_counted_as_a_rejected_upload(context):
    conn, _, actor = context
    digest = "d" * 64
    legacy = batches.create_batch(conn, source_file="old-side-workbook.xlsx",
                                  source_sheet="Old diagram", file_hash=digest,
                                  raw_store_key="wavescan/raw/" + digest,
                                  plant_code=PLANT, loaded_by=actor,
                                  template_fingerprint=None, raw_headers=None,
                                  sheet_class="OUT_OF_SCOPE_SHEET", domain="WAVESCAN")
    batches.set_batch_outcome(conn, legacy, status="REJECTED_TEMPLATE", notes="OUT_OF_SCOPE_SHEET: diagram")
    before = batches.count_batches(conn)
    data = accepted(context, workbook(ambiguous=True)).load_ids[0]
    assert batches.count_batches(conn) == before + 1
    assert legacy not in {row["load_id"] for row in batches.list_batches(conn, limit=200, offset=0)}
    assert batches.get_batch(conn, legacy)["status"] == "REJECTED_TEMPLATE"
    item = workbooks.get_workbook(conn, PLANT, digest)
    assert item["operational_summary"] == "SKIPPED"
    assert item["sheet_count"] == 0
    assert item["skipped_sheets"][0]["legacy_load_id"] == legacy
    assert conn.execute(text("SELECT rows_loaded+rows_rejected+rows_quarantine+rows_duplicate "
                             "FROM load_batch WHERE load_id=:id"), {"id": legacy}).scalar_one() == 0
    assert batches.get_batch(conn, data)["status"] == "REJECTED_TEMPLATE"

