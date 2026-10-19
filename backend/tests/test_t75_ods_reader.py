"""T75: ``.ods`` format gaps found by the real SAP FilmBuild dry run.

Every workbook here is built by hand in the test; no vault file is read.

1. **Time cells.** 13 of the 66 SAP files store the PELT ``Time:`` value as
   ``office:value-type="time"`` (``office:time-value="PT11H25M0S"``). The
   reader returned the display text; openpyxl returns ``datetime.time`` for
   the same cell in ``.xlsx``. Now both give ``datetime.time`` -- the time of
   day LibreOffice shows (``PT80H2M0S`` is shown and read as 08:02).
2. **Template JSON.** ``SheetTemplate.as_json()`` serialises a ``time`` in the
   PELT preamble, so ``create_batch``'s ``json.dumps`` cannot fail on it. A
   time-typed ``Time:`` cell in an ``.xlsx`` already hit that.
3. **Line breaks and tabs.** ``text:line-break`` / ``text:tab`` in a cell read
   as ``\\n`` / ``\\t`` (openpyxl's value for the same cell). No real file used
   them; latent.

Plus guards: a zip bomb and a DOCTYPE ``.ods`` are still refused, with the
limits unchanged.
"""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime, time

import pytest
from openpyxl import Workbook

from app.domain import ods
from app.domain.reader import describe_film_sheet, iter_rows
from app.services import upload_guard
from app.services.upload_guard import GuardLimits, UploadRejected, check_file

NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
)
FILM_ALIAS = {
    "location": "location", "clearcoat": "clear", "frozen white": "base",
    "white prime": "primer", "e-coat": "ecoat", "thickness": "total_thickness",
}


def text_cell(value: str) -> str:
    return f'<table:table-cell office:value-type="string"><text:p>{value}</text:p></table:table-cell>'


def float_cell(value: float) -> str:
    return (f'<table:table-cell office:value-type="float" office:value="{value}">'
            f"<text:p>{value}</text:p></table:table-cell>")


def blank_cell() -> str:
    return "<table:table-cell/>"


def build_ods(rows: list[list[str]], sheet: str = "QUICK_SUMMARY",
              prologue: str = "") -> bytes:
    """A minimal valid ``.ods`` from raw ``table:table-cell`` XML per row."""
    body = "".join(
        "<table:table-row>" + "".join(cells) + "</table:table-row>" for cells in rows
    )
    content = (
        f'<?xml version="1.0" encoding="UTF-8"?>{prologue}'
        f'<office:document-content {NS} office:version="1.2"><office:body>'
        f'<office:spreadsheet><table:table table:name="{sheet}">{body}'
        f'<table:table-row table:number-rows-repeated="1048000"><table:table-cell '
        f'table:number-columns-repeated="16384"/></table:table-row>'
        "</table:table></office:spreadsheet></office:body></office:document-content>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), ods.ODS_MIMETYPE)
        z.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("META-INF/manifest.xml", "<manifest/>", compress_type=zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def time_cell(duration: str, shown: str) -> str:
    return (f'<table:table-cell office:value-type="time" office:time-value="{duration}">'
            f"<text:p>{shown}</text:p></table:table-cell>")


def pelt_preamble(time_xml: str) -> list[list[str]]:
    """The SAP PELT Quick Summary shape: title, labels, header on row 14."""
    return [
        [text_cell("Axalta South Africa - PELT Filmbuild Report")],
        [],
        [text_cell("Date:"), text_cell("12-Aug-26")],
        [text_cell("Time:"), time_xml, blank_cell(), text_cell("Number of Points:"),
         float_cell(0)],
        [text_cell("Operator:"), text_cell("T75")],
        [text_cell("Job Number:"), text_cell("SYNTH")],
        [text_cell("Layer 1:"), text_cell("Clearcoat")],
        [text_cell("Layer 2:"), text_cell("Frozen White"), blank_cell(), text_cell("Vin:"),
         text_cell("SYN0001")],
        [text_cell("Layer 3:"), text_cell("White Prime")],
        [text_cell("Layer 4:"), text_cell("E-Coat")],
        [text_cell("Layer 5:"), text_cell("Steel")],
        [text_cell("Layer 6:")],
        [text_cell("Comments:")],
        [text_cell(h) for h in ("Location", "Clearcoat", "Frozen White", "White Prime",
                                "E-Coat", "Thickness")],
        [text_cell("001HOOD"), float_cell(40.5), float_cell(15.1), float_cell(25.2),
         float_cell(20.3), float_cell(101.1)],
    ]


# ---------------------------------------------------------------- 1. time
@pytest.mark.parametrize(
    "duration, expected",
    [
        ("PT11H25M0S", time(11, 25)),
        ("PT16H14M0S", time(16, 14)),
        ("PT1H16M0S", time(1, 16)),
        ("PT80H2M0S", time(8, 2)),          # real SAP file: shown as "8:02 am"
        ("PT09H40M05.5S", time(9, 40, 5, 500000)),
    ],
)
def test_a_time_cell_reads_as_datetime_time(duration, expected):
    data = build_ods([[text_cell("Time:"), time_cell(duration, "shown text")]])
    [row] = list(ods.iter_rows(data, "QUICK_SUMMARY"))
    assert row == ("Time:", expected)


def test_an_unreadable_time_value_keeps_its_text():
    data = build_ods([[text_cell("Time:"), time_cell("not-a-duration", "9:12 am")]])
    [row] = list(ods.iter_rows(data, "QUICK_SUMMARY"))
    assert row == ("Time:", "9:12 am")


def test_ods_and_xlsx_return_the_same_time_shape():
    data = build_ods([[text_cell("Time:"), time_cell("PT9H12M0S", "9:12 am")]])
    wb = Workbook()
    wb.active.title = "QUICK_SUMMARY"
    wb.active.append(["Time:", time(9, 12)])
    buffer = io.BytesIO()
    wb.save(buffer)
    from_ods = next(iter_rows(data, "QUICK_SUMMARY"))
    from_xlsx = next(iter_rows(buffer.getvalue(), "QUICK_SUMMARY"))
    assert from_ods == from_xlsx == ("Time:", time(9, 12))
    assert type(from_ods[1]) is type(from_xlsx[1])


# ---------------------------------------------------------------- 2. template JSON
def test_a_pelt_template_with_a_time_cell_serialises():
    data = build_ods(pelt_preamble(time_cell("PT11H25M0S", "11:25 am")))
    template = describe_film_sheet(data, "QUICK_SUMMARY", FILM_ALIAS)
    assert template is not None and template.header_row == 14
    assert template.meta_map["time"] == time(11, 25)
    as_json = template.as_json()
    assert as_json["meta"]["time"] == "11:25:00"
    json.dumps(as_json)  # create_batch dumps this without a default


def test_an_xlsx_pelt_template_with_a_time_cell_serialises():
    wb = Workbook()
    ws = wb.active
    ws.title = "QUICK_SUMMARY"
    ws.append(["Axalta South Africa - PELT Filmbuild Report"])
    ws.append([])
    ws.append(["Date:", datetime(2026, 8, 12)])
    ws.append(["Time:", time(23, 48), None, "Number of Points:", 0])
    for label, value in (("Operator:", "T75"), ("Job Number:", "SYNTH"),
                         ("Layer 1:", "Clearcoat"), ("Layer 2:", "Frozen White"),
                         ("Layer 3:", "White Prime"), ("Layer 4:", "E-Coat"),
                         ("Layer 5:", "Steel"), ("Layer 6:", ""), ("Comments:", "")):
        ws.append([label, value])
    ws.append(["Location", "Clearcoat", "Frozen White", "White Prime", "E-Coat",
               "Thickness"])
    ws.append(["001HOOD", 40.5, 15.1, 25.2, 20.3, 101.1])
    buffer = io.BytesIO()
    wb.save(buffer)
    template = describe_film_sheet(buffer.getvalue(), "QUICK_SUMMARY", FILM_ALIAS)
    assert template is not None and template.meta_map["time"] == time(23, 48)
    assert json.loads(json.dumps(template.as_json()))["meta"]["time"] == "23:48:00"


def test_ods_and_xlsx_pelt_rows_have_the_same_shape():
    data = build_ods(pelt_preamble(time_cell("PT11H25M0S", "11:25 am")))
    rows = list(iter_rows(data, "QUICK_SUMMARY"))
    assert len(rows) == 15                      # trailing 1,048,000 blank rows dropped
    assert rows[13] == ("Location", "Clearcoat", "Frozen White", "White Prime",
                        "E-Coat", "Thickness")
    assert rows[14] == ("001HOOD", 40.5, 15.1, 25.2, 20.3, 101.1)
    assert all(isinstance(r, tuple) for r in rows)


# ---------------------------------------------------------------- 3. line breaks / tabs
def test_line_break_and_tab_read_as_newline_and_tab():
    cell = ('<table:table-cell office:value-type="string"><text:p>Mid<text:line-break/>'
            'Grey<text:tab/>Prime<text:s text:c="2"/>x</text:p></table:table-cell>')
    data = build_ods([[cell]])
    [row] = list(ods.iter_rows(data, "QUICK_SUMMARY"))
    assert row == ("Mid\nGrey\tPrime  x",)


def test_a_cell_comment_is_not_cell_text():
    cell = ('<table:table-cell office:value-type="string"><office:annotation>'
            '<text:p>reviewer note</text:p></office:annotation>'
            '<text:p>001HOOD</text:p></table:table-cell>')
    data = build_ods([[cell]])
    assert list(ods.iter_rows(data, "QUICK_SUMMARY")) == [("001HOOD",)]


# ---------------------------------------------------------------- guards (unchanged)
def test_size_limits_are_unchanged():
    assert upload_guard.MAX_FILE_BYTES == 25 * 1024 * 1024
    assert upload_guard.MAX_UNCOMPRESSED_BYTES == 250 * 1024 * 1024
    assert upload_guard.MAX_COMPRESSION_RATIO == 200.0
    assert upload_guard.MAX_ARCHIVE_ENTRIES == 2000
    assert ".ods" in upload_guard.ALLOWED_EXTENSIONS


def test_an_ods_zip_bomb_is_still_refused():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), ods.ODS_MIMETYPE)
        z.writestr("content.xml", b"<a>" + b" " * (30 * 1024 * 1024) + b"</a>",
                   compress_type=zipfile.ZIP_DEFLATED)
    with pytest.raises(UploadRejected) as caught:
        check_file("bomb.ods", buffer.getvalue(), GuardLimits())
    assert caught.value.reason == "ZIP_BOMB"


def test_an_ods_with_a_doctype_is_still_refused():
    data = build_ods([[text_cell("x")]], prologue='<!DOCTYPE x [<!ENTITY a "b">]>')
    check_file("doctype.ods", data, GuardLimits())  # a small archive passes the door
    with pytest.raises(ods.OdsError, match="DOCTYPE"):
        ods.sheet_names(data)
    with pytest.raises(ods.OdsError, match="DOCTYPE"):
        list(ods.iter_rows(data, "QUICK_SUMMARY"))
