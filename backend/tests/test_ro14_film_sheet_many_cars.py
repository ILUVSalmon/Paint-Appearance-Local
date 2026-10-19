"""RO14 -- a Film sheet that holds several cars one after another loads as
several cars (every plant).

Claims:

* A sheet is split into walks at each row whose label resolves to the first
  point of the plant x model's live list (D-RO14-1); seven one-car walks are
  seven cars, and two cars of one colour on one day with no Job Number are
  two cars with two ids (D-RO14-4).
* Each walk is checked on its own; a car out of order holds the whole sheet
  with a sentence naming that car (D-RO14-2), and a walk carrying two
  colours holds it too (D-RO14-3). The tab reads that hold's own sentence,
  not "waiting for an Admin" (D-RO14-6), and the sentence has no ``**``
  (D-RO14-7).
* A one-walk sheet keeps every id it had, and every stored Film car's id
  still recomputes the same; a renamed copy of a committed multi-walk sheet
  is all ``DUPLICATE``; a plant with no Film settings still splits for
  identity but checks no order.

Synthetic workbooks only. Every batch a test makes is removed by
``film_cleanup`` (facts first); no reference row is written.
"""

from __future__ import annotations

import hashlib
import io
import uuid

import pytest
from openpyxl import Workbook

from app.domain.film import logical_vehicle_id
from app.domain.film_walk import check_walk, check_walks, walk_fingerprint, walk_starts
from tests.test_t04_ingest import forget_bytes, unique_rows
from tests.test_t04_ingest import admin, cleanup, engineer, ingest_apps, viewer  # noqa: F401
from tests.test_t18_profiles import ingest_as, signature_of
from tests.test_t18_profiles import owned_profile, reordered  # noqa: F401  (fixtures)
from tests.test_t18_profiles import workbook as t18_workbook
from tests.test_t46_filmthickness import film_cleanup, sent  # noqa: F401
from tests.test_t54_filmthickness_polish import FTM_GOLD, desert_island, old_day, ws_url
from tests.ro19_synthetic import ro19_plants  # noqa: F401

pytestmark = pytest.mark.ingest

ADMIN_WAIT = "This tab cannot load yet: its layout is waiting for an Admin to approve it."
MODEL = "P703 DBL RANGER"
WALK = ["001HOOD", "002HOOD", "003HOOD"]
#: Seven colours AAT offers (`ref_plant_color`), as their catalogue names.
SEVEN = ["Acacia Green", "Aluminum Metallic", "Arctic White", "Command Grey",
         "Ignite Orange", "Meteor Grey", "Traction Green"]
#: The owner's July AAT layout, with the Date column the owner added.
HEADERS = ["Checkzone", "Clearcoat", "Basecoat2", "Basecoat1", "E-coat",
           "Orientation", "Color", "Date"]


def a_day() -> str:
    """A day no real car was measured on, as the Date column writes it."""
    n = uuid.uuid4().int
    return f"20{1 + n % 15:02d}-{1 + (n // 15) % 12:02d}-{1 + (n // 180) % 28:02d}"


def cars_sheet(cars, *, tab: str = "P703") -> bytes:
    """One tab, cars one after another. ``cars`` is ``[(labels, colours, day)]``;
    ``colours`` is one name or one per label. Readings differ per car."""
    wb = Workbook()
    wb.properties.title = uuid.uuid4().hex  # distinct bytes for equal readings
    ws = wb.active
    ws.title = tab
    ws.append(HEADERS)
    for k, (labels, colours, day) in enumerate(cars):
        names = [colours] * len(labels) if isinstance(colours, str) else colours
        for i, (label, colour) in enumerate(zip(labels, names)):
            ws.append([label, 40.0 + i, 10.0 + k, 8.0 + k / 10, 20.0, "H", colour, day])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def renamed(data: bytes) -> bytes:
    """The same readings in other bytes (a renamed, re-saved copy)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    wb.properties.title = uuid.uuid4().hex
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def only_tab(engineer, plant: str, review: dict) -> dict:  # noqa: F811
    body = engineer.get(ws_url(plant, review["batch"]["file_hash"])).json()
    [tab] = [s for s in body["sheets"] if s["load_id"] == review["batch"]["load_id"]]
    return tab


def review_rows(admin_conn, load_id: int) -> list[tuple]:
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT source_row, disposition, reason, detail, "
            "       canonical->>'logical_vehicle_id' "
            "  FROM stg_review_row WHERE load_id = %s ORDER BY source_row",
            (load_id,),
        )
        return cur.fetchall()


def upload_one(engineer, film_cleanup, admin_conn, name, data, plant):  # noqa: F811
    [review] = sent(engineer, film_cleanup, admin_conn, name, data, plant).values()
    return review


# =====================================================================
# A. pure rules
# =====================================================================
def test_a_sheet_splits_at_the_first_point_and_one_walk_is_check_walk():
    labels = WALK + WALK + ["1)001HOOD", "002HOOD"]
    assert walk_starts(labels, WALK + ["004HOOD"]) == [0, 3, 6]
    # An alias of the first point starts a car too; no live list, no split.
    assert walk_starts(["HOOD A", "002HOOD", "HOOD A"], WALK, {"HOOD A": "001HOOD"}) == [0, 2]
    assert walk_starts(labels, []) == [0]
    # One Job Number on every row (a PELT preamble) names one car: a first
    # point written again is the walk guard's repeat, never a second car.
    assert walk_starts(labels, WALK, identities=[("J1", None)] * len(labels)) == [0]
    assert walk_starts(labels, WALK, identities=[(None, None)] * len(labels)) == [0, 3, 6]
    twice = ["001HOOD", "002HOOD", "001HOOD", "003HOOD"]
    held = check_walks(twice, WALK, rows=[2, 3, 4, 5], plant="X", model="M",
                       identities=[("J1", None)] * 4)
    assert held.holds and held.detail.startswith('This sheet does not walk')
    # One walk is exactly today's verdict, sentence included.
    one = ["001HOOD", "003HOOD", "002HOOD"]
    assert check_walks(one, WALK, rows=[2, 3, 4], plant="X", model="M") == check_walk(
        one, WALK, plant="X", model="M")


def test_the_walk_sentence_names_the_car_and_has_no_markdown():
    labels = WALK + WALK + ["001HOOD", "003HOOD", "002HOOD"]
    verdict = check_walks(labels, WALK, rows=list(range(2, 11)), plant="X", model="M")
    assert verdict.holds and not verdict.ok
    assert verdict.detail.startswith(
        "Car 3 of 3 on this sheet (starting at row 8): This sheet does not walk "
        "X's M walk in its own order")
    assert "**" not in verdict.detail
    assert "**" not in check_walk(["002HOOD", "001HOOD"], WALK, plant="X", model="M").detail
    ok = check_walks(WALK + WALK[:2], WALK, rows=list(range(2, 7)), plant="X", model="M")
    assert ok.ok and not ok.holds


def test_the_fingerprint_reads_points_and_layers_not_the_file():
    a = [("001HOOD", (40.0, 10.0, 8.0, 20.0)), ("002HOOD", (41.0, 10.0, 8.0, None))]
    assert walk_fingerprint(a) == walk_fingerprint(list(a))
    assert len(walk_fingerprint(a)) == 16
    assert walk_fingerprint(a) != walk_fingerprint([a[0], ("002HOOD", (41.0, 10.0, 8.0, 0))])


def test_every_stored_film_car_keeps_its_id(admin_conn):
    """D-RO14-4: the optional argument left out, every id is what it was."""
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT DISTINCT plant, measured_date, model, color, job_no, serial_no, "
            "       logical_vehicle_id FROM filmthickness"
        )
        stored = cur.fetchall()
    assert stored, "the standing database has Film facts to compare"
    assert [r for r in stored if logical_vehicle_id(*r[:6]) != r[6]] == []
    # A known one, spelt out: today's formula, no walk in it.
    plain = hashlib.sha256(b"AAT|2026-07-13|P703 DBL RANGER|Command Grey|blank").hexdigest()
    from datetime import date
    assert logical_vehicle_id(
        "AAT", date(2026, 7, 13), MODEL, "Command Grey", None, None) == plain
    assert logical_vehicle_id(
        "AAT", date(2026, 7, 13), MODEL, "Command Grey", "J1", None, "abc") == (
        logical_vehicle_id("AAT", date(2026, 7, 13), MODEL, "Command Grey", "J1", None))


# =====================================================================
# B. through the app (AAT has Film settings)
# =====================================================================
def test_seven_one_car_walks_with_seven_colours_are_seven_cars(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    data = cars_sheet([(WALK, colour, day) for colour in SEVEN])
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 seven {uuid.uuid4().hex[:6]}.xlsx",
                        data, "AAT")
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert {r[1] for r in rows} == {"LOAD"} and len(rows) == 21
    assert len({r[4] for r in rows}) == 7
    tab = only_tab(engineer, "AAT", review)
    assert (tab["state"], tab["vehicle_count"], tab["rows_ready"]) == ("ready", 7, 21)


def test_two_cars_of_one_colour_on_one_day_with_no_job_are_two_cars(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    data = cars_sheet([(WALK, "Command Grey", day), (WALK, "Command Grey", day)])
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 twins {uuid.uuid4().hex[:6]}.xlsx",
                        data, "AAT")
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert [r[1] for r in rows] == ["LOAD"] * 6
    first, second = {r[4] for r in rows[:3]}, {r[4] for r in rows[3:]}
    assert len(first) == len(second) == 1 and first != second
    # Neither is the one id the old rule gave both.
    from datetime import date
    folded = logical_vehicle_id("AAT", date.fromisoformat(day), MODEL, "Command Grey", None, None)
    assert folded not in first | second
    assert only_tab(engineer, "AAT", review)["vehicle_count"] == 2


def test_a_car_out_of_order_holds_the_sheet_and_the_tab_says_which_car(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    cars = [(WALK, "Command Grey", day), (WALK, "Meteor Grey", day),
            (["001HOOD", "003HOOD", "002HOOD"], "Arctic White", day),
            (WALK, "Acacia Green", day)]
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 order {uuid.uuid4().hex[:6]}.xlsx",
                        cars_sheet(cars), "AAT")
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert {(r[1], r[2]) for r in rows} == {("QUARANTINE", "FILM_WALK_UNKNOWN")}
    detail = rows[0][3]
    assert detail.startswith("Car 3 of 4 on this sheet (starting at row 8): ")
    assert '"002HOOD" after "003HOOD"' in detail and "**" not in detail
    tab = only_tab(engineer, "AAT", review)
    assert (tab["state"], tab["operator_message"]) == ("blocked", detail)
    assert "Admin" not in tab["operator_message"]


def test_a_car_carrying_two_colours_holds_the_sheet(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    cars = [(WALK, "Command Grey", day),
            (WALK, ["Meteor Grey", "Meteor Grey", "Arctic White"], day)]
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 two colours {uuid.uuid4().hex[:6]}.xlsx",
                        cars_sheet(cars), "AAT")
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert {(r[1], r[2]) for r in rows} == {("QUARANTINE", "FILM_WALK_UNKNOWN")}
    assert rows[0][3] == (
        'Car 2 of 2 on this sheet (starting at row 5) carries two colours, '
        '"Meteor Grey" and "Arctic White". One car is one day and one colour, so '
        "the whole sheet is held and nothing from it loads. Fix the file and "
        "upload it again."
    )
    assert only_tab(engineer, "AAT", review)["operator_message"] == rows[0][3]


def test_a_one_walk_sheet_keeps_its_ids_and_answers(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 one {uuid.uuid4().hex[:6]}.xlsx",
                        cars_sheet([(WALK, "Command Grey", day)]), "AAT")
    plain = hashlib.sha256(
        f"AAT|{day}|{MODEL}|Command Grey|blank".encode()).hexdigest()
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert [(r[1], r[2], r[4]) for r in rows] == [("LOAD", None, plain)] * 3
    tab = only_tab(engineer, "AAT", review)
    assert (tab["state"], tab["vehicle_count"]) == ("ready", 1)

    # D-RO14-11: an FTM-shaped one-car PELT, Job Number in its preamble.
    ftm_day, _ = old_day()
    job = f"IB_RB_{uuid.uuid4().int % 10**7:07d}"
    pelt = sent(engineer, film_cleanup, admin_conn, f"RO14 {job}.xlsx",
                desert_island(FTM_GOLD[:3], day=ftm_day, job=job), "FTM")["Quick Summary"]
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT DISTINCT canonical->>'measured_date', canonical->>'model', "
            "       canonical->>'color', canonical->>'logical_vehicle_id' "
            "  FROM stg_review_row WHERE load_id = %s AND disposition = 'LOAD'",
            (pelt["batch"]["load_id"],),
        )
        [(measured, model, colour, vid)] = cur.fetchall()
    assert vid == hashlib.sha256(
        f"FTM|{measured}|{model}|{colour}|job:{job}".encode()).hexdigest()
    assert only_tab(engineer, "FTM", pelt)["state"] == "ready"


def test_a_renamed_copy_of_a_committed_multi_walk_sheet_is_all_duplicate(
    engineer, admin_conn, film_cleanup  # noqa: F811
):
    day = a_day()
    data = cars_sheet([(WALK, "Command Grey", day), (WALK, "Command Grey", day)])
    first = upload_one(engineer, film_cleanup, admin_conn, f"RO14 first {uuid.uuid4().hex[:6]}.xlsx",
                       data, "AAT")
    loaded = engineer.post(ws_url("AAT", first["batch"]["file_hash"], "load-ready"))
    assert loaded.status_code == 200, loaded.text
    copy = renamed(data)
    assert copy != data
    again = upload_one(engineer, film_cleanup, admin_conn, f"RO14 copy {uuid.uuid4().hex[:6]}.xlsx",
                       copy, "AAT")
    rows = review_rows(admin_conn, again["batch"]["load_id"])
    assert [r[1] for r in rows] == ["DUPLICATE"] * 6
    assert len({r[4] for r in rows}) == 2


# =====================================================================
# C. a plant with no Film settings (FVL): split for identity, no order check
# =====================================================================
def test_a_plant_without_settings_splits_for_identity_but_checks_no_order(
    engineer, admin_conn, film_cleanup, ro19_plants  # noqa: F811
):
    plant = ro19_plants()
    day = a_day()
    cars = [(WALK, "Pace Red", day), (["001HOOD", "003HOOD", "002HOOD"], "Pace Red", day)]
    review = upload_one(engineer, film_cleanup, admin_conn, f"RO14 FVL {uuid.uuid4().hex[:6]}.xlsx",
                        cars_sheet(cars), plant)
    rows = review_rows(admin_conn, review["batch"]["load_id"])
    assert [r[1] for r in rows] == ["LOAD"] * 6, rows
    assert len({r[4] for r in rows[:3]} | {r[4] for r in rows[3:]}) == 2


# =====================================================================
# D. D-RO14-6: a layout really waiting on an Admin still says so
# =====================================================================
def test_a_real_profile_wait_still_reads_the_admin_sentence(
    engineer, admin_conn, cleanup, owned_profile, reordered  # noqa: F811
):
    data = t18_workbook(unique_rows(3), order=reordered)
    forget_bytes(admin_conn, data)
    owned_profile("AAT", signature_of(admin_conn, data), fixture_verified=False)
    review = ingest_as(engineer, cleanup, "ro14-unverified.xlsx", data, "AAT")
    assert review["commit_blockers"] == ["UNVERIFIED_PROFILE"]
    tab = only_tab(engineer, "AAT", review)
    assert (tab["state"], tab["operator_message"]) == ("blocked", ADMIN_WAIT)
