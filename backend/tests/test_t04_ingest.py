"""T04 -- ingest, against the real database and the real worker.

Everything here runs the actual FastAPI application, the actual PostgreSQL
schema, and the actual background worker claiming jobs with `SKIP LOCKED`.
The claims T04 has to make are claims about a running system -- that a
hostile file leaves no batch, that identical bytes are refused by a
constraint rather than by an `if`, that a committed batch appears in the
measurement list and an uncommitted one does not -- and none of those can be
demonstrated against a mock.

Where the vault's own workbooks are reachable they are used, because a
synthetic file proves the parser handles synthetic files. The real AAT June
2026 workbook is 629 rows of Ford plant data and is **never copied into this
repository** (`03-SECURITY.md` §8): it is read from the vault at test time
and skipped if the vault is absent.

Every test cleans up after itself. It has to: the reload guard means a
second run would otherwise be refused as a duplicate upload, which would
look like a bug and would in fact be the guard working.

T17 CHANGE: every upload here now sends `expected_plant=AAT`, because the
route requires it. The plant gate itself -- wrong plant, mixed sheets, blank
cells filled from the declared plant -- is tested in
`tests/test_t17_expected_plant.py`.
"""

from __future__ import annotations

import io
import os
import time
import uuid
from datetime import date
from pathlib import Path

import psycopg
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.config import Settings
from app.domain.cleaner import MONTH_NAMES
from app.main import create_app
from app.repositories import ingest_jobs
from app.workers import ingest_worker

pytestmark = pytest.mark.ingest

BACKEND_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND_ROOT.parent

#: Read-only reference vault (`00-ARCHITECTURE.md` §8). Never written to.
VAULT = Path(
    os.environ.get(
        "WAVESCAN_VAULT_XLSX",
        str(REPO_ROOT.parent / "Paint Appearance" / "Component" / "Database"
            / "Excel Template" / "WaveScan"),
    )
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADERS = [
    "Date", "Month", "Year", "Plant", "Model", "Checkzone", "Color",
    "Tolerance Group", "Paintline", "Serial No",
    "LW", "SW", "CF", "DOI", "du", "Wa", "Wb", "Wc", "Wd", "We", "B",
]

#: Values that resolve through the seeded alias tables. A U704 has RRHOOD.
BASE_ROW = {
    "Date": "2026-08-03", "Month": "August", "Year": 2026,
    "Plant": "AAT", "Model": "U704", "Checkzone": "01 RRHOOD",
    "Color": "3HC Traction Green", "Tolerance Group": "Horizontal",
    "Paintline": "BOOTH1", "Serial No": "T04-0001",
    "LW": 3.7, "SW": 15.7, "CF": 60.9, "DOI": 87.5, "du": 15.9,
    "Wa": 13.4, "Wb": 21.5, "Wc": 11.6, "Wd": 12.7, "We": 6.7, "B": -0.2,
}

#: Signed offered ticks (STANDARDS-D2-AUDIT). unique_rows that only override
#: Plant must still load; missing key on the live maps is not-offered.
OFFERED_FIXTURE = {
    "AAT": {"Model": "U704", "Color": "3HC Traction Green"},
    "FTM": {"Model": "P703 DBL RANGER", "Color": "Ignite Orange"},
    "FVL": {"Model": "P703 DBL RANGER", "Color": "Ignite Orange"},
    "SAP": {"Model": "P703 DBL RANGER", "Color": "Frozen White"},
}


#: The first five panels of each plant's WaveScan gold (T58).
PLANT_ZONES = {
    "AAT": ["01 RRHOOD", "02 RFHOOD", "03 CFHOOD", "04 LFHOOD", "05 LRHOOD"],
    "FTM": ["01 RRHOOD", "02 RFHOOD", "03 CFHOOD", "04 LFHOOD", "05 LRHOOD"],
    "FVL": ["1HOOD", "2HOOD", "3HOOD", "4HOOD", "5HOOD"],
    "SAP": ["01RHS Hood Rr28", "02RHS Hood Frnt29", "03Hood Middle30",
            "04LHS Hood Frnt31", "05LHS Hood Rr32"],
}


def on_base_row(row: dict) -> dict:
    """``BASE_ROW`` under ``row``. RO15: a Date the row sets carries its own
    Month / Year, as a real export does, unless the row sets those too --
    a row whose date disagrees with them is held."""
    merged = {**BASE_ROW, **row}
    try:
        day = date.fromisoformat(str(row["Date"])[:10])
    except (KeyError, ValueError):
        return merged
    return {**merged, "Month": row.get("Month", MONTH_NAMES[day.month - 1]),
            "Year": row.get("Year", day.year)}


def make_workbook(rows: list[dict], sheets: dict[str, list[dict]] | None = None) -> bytes:
    """A workbook whose values resolve against the seeded reference data."""
    wb = Workbook()
    sheets = sheets or {"PasteData": rows}
    first = True
    for name, sheet_rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        ws.append(HEADERS)
        for row in sheet_rows:
            merged = on_base_row(row)
            ws.append([merged[h] for h in HEADERS])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def unique_rows(count: int = 3, **over) -> list[dict]:
    """Rows for a vehicle no other test has used, so dedup stays out of it."""
    serial = f"T04-{uuid.uuid4().hex[:10]}"
    plant = over.get("Plant", BASE_ROW["Plant"])
    # T58 CHANGE (D-58-01): FVL and SAP keep their own panel names; AAT/FTM
    # long-form is not their gold, so their rows write their own walk.
    zones = PLANT_ZONES.get(plant, PLANT_ZONES["AAT"])
    identity = OFFERED_FIXTURE.get(plant, OFFERED_FIXTURE["AAT"])
    return [
        {"Serial No": serial, "Checkzone": zones[i % len(zones)],
         "CF": 60.0 + i, **identity, **over}
        for i in range(count)
    ]


# =====================================================================
# fixtures
# =====================================================================
def _settings(url: str, **overrides) -> Settings:
    base = dict(
        env="development",
        app_database_url=url,
        auth_dev_bypass=True,
        auth_dev_bypass_roles_raw="engineer",
        # The default store, shared with every other client in the session:
        # a worker in one app must be able to read bytes another app stored.
        ingest_worker_poll_seconds=0.05,
    )
    base.update(overrides)
    return Settings(**base)


# T76 CHANGE: module scope, not session. Every test module that imports this
# fixture gets its own copy (pytest registers the import per module), so at
# session scope ~40 modules each kept their apps and connection pools open
# until the run ended -- past the cluster's 60 connections. Module scope
# tears each copy down (and disposes its engine) when its module is done.
@pytest.fixture(scope="module")
def ingest_apps(app_database_url: str):
    created = []

    def _make(**overrides) -> TestClient:
        app = create_app(_settings(app_database_url, **overrides))
        client = TestClient(app)
        client.__enter__()
        created.append((client, app))
        return client

    yield _make
    for client, app in created:
        client.__exit__(None, None, None)
        app.state.database.dispose()


@pytest.fixture(scope="module")
def engineer(ingest_apps) -> TestClient:
    return ingest_apps()


@pytest.fixture(scope="module")
def viewer(ingest_apps) -> TestClient:
    return ingest_apps(auth_dev_bypass_roles_raw="viewer", ingest_worker_enabled=False)


@pytest.fixture(scope="module")
def admin(ingest_apps) -> TestClient:
    """An Admin caller, with no worker of its own.

    T19 needs one here rather than only in its own module: acknowledging a
    warning is Admin, and from T19 a sheet with rejected, quarantined or
    duplicate rows cannot be committed until somebody has. The parsing is
    still done by the engineer client's worker -- a second worker would race
    it for jobs and prove nothing.
    """
    return ingest_apps(auth_dev_bypass_roles_raw="admin", ingest_worker_enabled=False)


def _wait_for_idle_jobs(admin_conn, load_ids: list[int], timeout: float = 20.0) -> None:
    """Block until no ingest job for these batches is QUEUED or RUNNING.

    T88: the same worker race T87 (D-87-8) fixed in `film_cleanup` is in
    this fixture too, and it leaked again -- `t86-panel-match.xlsx` was left
    standing in the owner's live database as `load_id 63055`, with
    `ingest_job` gone and 5 `stg_raw` / 5 `stg_review_row` rows re-created
    after the teardown had deleted them. The closing `DELETE FROM
    load_batch` then hit `ON DELETE RESTRICT` and the batch survived.

    Deleting the job row does not stop a worker that has already claimed
    it, so the only correct move is to wait for it to finish first.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        with admin_conn.cursor() as cur:
            cur.execute(
                "SELECT count(*) FROM ingest_job "
                " WHERE load_id = ANY(%s) AND status IN ('QUEUED', 'RUNNING')",
                (load_ids,),
            )
            if cur.fetchone()[0] == 0:
                return
        time.sleep(0.1)


@pytest.fixture()
def cleanup(admin_conn):
    """Remove every batch a test created, in FK order.

    The application cannot do this -- no role holds DELETE on a fact table,
    which is T01's whole result -- so it runs as the owner, and only for rows
    this suite created.

    T88: it waits for the ingest worker, retries once, and then **fails the
    run** if a batch is still standing -- so a leak into the owner's live
    database becomes a visible test error instead of silent residue
    somebody finds months later. Same shape as `film_cleanup` since T87.
    """
    with admin_conn.cursor() as cur:
        cur.execute("SELECT plant_code,file_hash FROM workbook_upload")
        workbook_keys_before = set(cur.fetchall())
    load_ids: list[int] = []
    yield load_ids
    leaked: list[int] = []
    for _ in (1, 2):
        if load_ids:
            _wait_for_idle_jobs(admin_conn, load_ids)
        with admin_conn.cursor() as cur:
            if load_ids:
                _forget_batches(cur, load_ids)
                cur.execute(
                    "SELECT load_id FROM load_batch WHERE load_id = ANY(%s) ORDER BY 1",
                    (load_ids,),
                )
                leaked = [row[0] for row in cur.fetchall()]
            cur.execute("SELECT plant_code,file_hash FROM workbook_upload")
            new_workbooks = set(cur.fetchall()) - workbook_keys_before
            for plant, digest in new_workbooks:
                cur.execute(
                    "DELETE FROM audit_event WHERE entity_type='workbook_upload' "
                    "AND entity_id=%s AND NOT EXISTS (SELECT 1 FROM load_batch "
                    "WHERE plant_code=%s AND file_hash=%s)",
                    (f"{plant}:{str(digest).strip()}", plant, digest),
                )
                cur.execute(
                    "DELETE FROM workbook_upload WHERE plant_code=%s AND file_hash=%s "
                    "AND NOT EXISTS (SELECT 1 FROM load_batch WHERE plant_code=%s AND file_hash=%s)",
                    (plant, digest, plant, digest),
                )
            cur.execute("SELECT plant_code,file_hash FROM workbook_upload")
            leaked_workbooks = set(cur.fetchall()) - workbook_keys_before
        if not leaked and not leaked_workbooks:
            return
    raise AssertionError(
        f"cleanup could not remove batches {leaked} or workbook evidence "
        f"{sorted(leaked_workbooks)} after two attempts. Those rows "
        f"are now residue in the standing database and must be removed by "
        f"load_id (T87 D-87-8, seen again in T88)."
    )


def forget_proposals(
    admin_conn, before: set[str], markers: list[str], who: str
) -> None:
    """Remove the mapping proposals one test created, or fail the run.

    **T92 Part 5**, the leak T89, T90 and T91 each reported and none owned:
    440 -> 477 `proposed_mapping` rows across T91's three suite runs. The
    modules that create proposals delete the *aliases* they wrote -- T29's
    `zone_cleanup` is careful about `ref_plant_zone_map` -- and then leave
    the proposal itself behind, because no fixture ever owned it. Three
    reports in, 513 rows of residue sit in the owner's live trial database.

    Same shape as `cleanup` and T87's `film_cleanup` (D-87-8): delete by
    **id**, read back, retry once, then **fail loudly** rather than leak.
    A test that cannot clean up after itself is a test failure; silent
    residue is the thing that goes unnoticed for months.

    Two conditions, both required, which is the T76 shape `row_keys` /
    `delete_rows_created_since` exists for: the row was **not** there before
    this test ran, **and** its payload carries one of `markers`, a substring
    this module writes. T30 and T31 post the owner's own G1 pack with the
    owner's own name on it, and a cleanup that deleted by content alone
    would erase the record of how the live owner-signed rows came to be. A
    row that was already there is never touched, whoever wrote it.
    """
    patterns = [f"%{marker}%" for marker in markers]
    mine = (
        "SELECT id FROM proposed_mapping"
        " WHERE payload::text LIKE ANY(%s)"
        "   AND NOT (id::text = ANY(%s))"
    )
    args = (patterns, list(before))
    leaked: list[str] = []
    for _ in (1, 2):
        with admin_conn.cursor() as cur:
            cur.execute(mine, args)
            ids = [row[0] for row in cur.fetchall()]
            if not ids:
                return
            cur.execute("DELETE FROM proposed_mapping WHERE id = ANY(%s)", (ids,))
            cur.execute(mine, args)
            leaked = [str(row[0]) for row in cur.fetchall()]
        if not leaked:
            return
    raise AssertionError(
        f"{who} could not remove its mapping proposals {leaked} after two "
        f"attempts. Those rows are now residue in the standing database and "
        f"must be removed by id (T92 Part 5, the leak T89/T90/T91 reported)."
    )


def _forget_batches(cur, load_ids: list[int]) -> None:
    """Remove every trace of these batches, in foreign-key order.

    T09 CHANGE: `measurement_supersession` is deleted first. It references
    `wavescan (measurement_id)` with `ON DELETE RESTRICT`, so once a test has
    corrected a measurement the old order failed with a RestrictViolation --
    which is the constraint doing its job. A retirement cannot be erased
    while the measurement it retires still exists, by anyone, which is why
    this runs as the owner on a throwaway database and never as `app_rw`.
    """
    cur.execute("SELECT DISTINCT plant_code,file_hash FROM load_batch WHERE load_id=ANY(%s)", (load_ids,))
    owned_workbooks = cur.fetchall()
    cur.execute(
        """
        DELETE FROM measurement_supersession
         WHERE superseded_measurement_id IN (
                   SELECT measurement_id FROM wavescan WHERE load_id = ANY(%s))
            OR superseded_by_measurement_id IN (
                   SELECT measurement_id FROM wavescan WHERE load_id = ANY(%s))
        """,
        (load_ids, load_ids),
    )
    # A correction points at the row it corrects, so the children go first.
    cur.execute(
        "DELETE FROM wavescan WHERE load_id = ANY(%s) "
        "AND corrects_measurement_id IS NOT NULL",
        (load_ids,),
    )
    # T18 CHANGE: `stg_review_row` joins the list. It references load_batch
    # with ON DELETE RESTRICT, like every other child, so a batch cannot be
    # forgotten while its projection still describes it.
    # T19 CHANGE: `review_acknowledgement` joins the list, for the same
    # reason and with the same force: it references load_batch ON DELETE
    # RESTRICT, and the application holds no DELETE on it at all.
    # T21 CHANGE: `review_override` joins it, for the same reason.
    for table in ("ingest_job", "stg_raw", "reject_log", "stg_review_row",
                  "review_acknowledgement", "review_override", "unidentified_model",
                  "proposed_mapping", "wavescan"):
        cur.execute(f"DELETE FROM {table} WHERE load_id = ANY(%s)", (load_ids,))
    cur.execute("DELETE FROM load_batch WHERE load_id = ANY(%s)", (load_ids,))
    # RO19: erase only metadata whose last owned test batch was removed.
    for plant, digest in owned_workbooks:
        cur.execute("DELETE FROM audit_event a WHERE a.entity_type='workbook_upload' "
                    "AND a.entity_id=%s AND NOT EXISTS (SELECT 1 FROM load_batch b "
                    "WHERE b.plant_code=%s AND b.file_hash=%s)",
                    (f"{plant}:{str(digest).strip()}", plant, digest))
        cur.execute("DELETE FROM workbook_upload u WHERE plant_code=%s AND file_hash=%s "
                    "AND NOT EXISTS (SELECT 1 FROM load_batch b WHERE b.plant_code=u.plant_code "
                    "AND b.file_hash=u.file_hash)", (plant, digest))


def forget_bytes(admin_conn, data: bytes) -> None:
    """Remove any batch already holding these exact bytes, as the OWNER.

    The reload guard makes an upload of the same file a 409, which is
    correct and is asserted elsewhere. It also means a test that ingests a
    vault workbook cannot run twice -- and the local harness ingests the
    same workbooks. So this suite clears its own precondition rather than
    depending on the order somebody ran things in.

    The application role cannot do this and never will (T01, ADR 0005); the
    owner can, and this is a throwaway test database.
    """
    import hashlib

    digest = hashlib.sha256(data).hexdigest()
    with admin_conn.cursor() as cur:
        cur.execute("SELECT load_id FROM load_batch WHERE file_hash = %s", (digest,))
        stale = [r[0] for r in cur.fetchall()]
        if stale:
            _forget_batches(cur, stale)
        # RO19: a purge, or a plant answer that frees the hash, deletes the
        # batch and leaves the `workbook_upload` row -- so no batch then
        # carries the hash into `_forget_batches` and the row outlives the
        # test. These are this suite's own bytes, so clear their evidence too.
        cur.execute("SELECT plant_code FROM workbook_upload WHERE file_hash = %s", (digest,))
        for (plant,) in cur.fetchall():
            cur.execute("DELETE FROM audit_event WHERE entity_type = 'workbook_upload' "
                        "AND entity_id = %s", (f"{plant}:{digest}",))
        cur.execute("DELETE FROM workbook_upload WHERE file_hash = %s", (digest,))


def upload(
    client: TestClient, name: str, data: bytes, expected_plant: str | None = "AAT"
):
    """POST one workbook.

    T17 CHANGE: `expected_plant` is a required multipart field. It defaults to
    AAT here because `BASE_ROW` is an AAT row, so every test that was written
    before the field existed keeps meaning what it meant. Pass
    `expected_plant=None` to send the pre-T17 body and prove the route refuses
    it; pass another code to exercise the plant gate.
    """
    body = {"files": (name, data, XLSX_MIME)}
    if expected_plant is None:
        return client.post("/api/v1/uploads", files=body)
    return client.post(
        "/api/v1/uploads", files=body, data={"expected_plant": expected_plant}
    )


def wait_for(client: TestClient, load_id: int, timeout: float = 30.0) -> dict:
    """Poll until the worker has finished with this batch."""
    deadline = time.monotonic() + timeout
    body = {}
    while time.monotonic() < deadline:
        body = client.get(f"/api/v1/batches/{load_id}").json()
        if body["batch"]["status"] != "RUNNING":
            return body
        time.sleep(0.05)
    raise AssertionError(f"batch {load_id} still RUNNING after {timeout}s: {body}")


def acknowledge(admin_client: TestClient, review: dict, category: str,
                note: str = "Reviewed by the T04 suite."):
    """Accept one warning category so a sheet with warnings can be committed.

    T19 CHANGE. A sheet whose parse rejected, quarantined or duplicated rows
    is not committable until an Admin has said, in writing and against this
    exact review revision, that they read them. Tests that commit such a
    sheet therefore have to do what a person would do -- which is the point
    of the gate, and why they were updated rather than exempted.
    """
    batch = review["batch"]
    return admin_client.post(
        f"/api/v1/plants/{batch['plant_code']}/workbooks/"
        f"{batch['file_hash']}/acknowledgements",
        json={
            "load_id": batch["load_id"],
            "category": category,
            "review_revision": batch["review_revision"],
            "note": note,
        },
    )


#: Which count on a batch makes which acknowledgement mandatory (T19).
REQUIRED_ACK_CATEGORIES = (
    ("rows_rejected", "REJECTED_ROWS"),
    ("rows_quarantine", "QUARANTINED_MODELS"),
    ("rows_duplicate", "DUPLICATE_ROWS"),
)


def acknowledge_required(admin_client: TestClient, review: dict) -> list[str]:
    """Acknowledge every category this batch's counts make mandatory.

    A real workbook's mix of rejects and duplicates is data-dependent, so a
    test that hardcoded the list would go green or red depending on what else
    the database already holds. This asks the batch.
    """
    done: list[str] = []
    counts = review["batch"]["counts"]
    for column, category in REQUIRED_ACK_CATEGORIES:
        if not counts[column]:
            continue
        response = acknowledge(admin_client, review, category)
        assert response.status_code == 201, response.text
        done.append(category)
    return done


def ingest(client, cleanup, name: str, data: bytes) -> dict:
    """Upload one file, wait for its single sheet, return the review body."""
    response = upload(client, name, data)
    assert response.status_code == 202, response.text
    load_id = response.json()["batches"][0]["load_id"]
    cleanup.append(load_id)
    return wait_for(client, load_id)


# =====================================================================
# 1. the happy path
# =====================================================================
def test_an_engineer_can_upload_and_the_worker_parses_it(engineer, cleanup):
    review = ingest(engineer, cleanup, "t04-happy.xlsx", make_workbook(unique_rows(3)))
    batch = review["batch"]
    assert batch["status"] == "PENDING_REVIEW"
    assert batch["counts"] == {
        "rows_in_file": 3, "rows_loaded": 3, "rows_duplicate": 0,
        "rows_rejected": 0, "rows_quarantine": 0, "balanced": True,
    }
    assert batch["job"]["status"] == "DONE"
    assert batch["plant_code"] == "AAT"
    assert review["staged_rows"] == 3, "stg_raw holds the file verbatim"


def test_nothing_reaches_wavescan_before_commit(engineer, cleanup, admin_conn):
    """`02-PIPELINE.md` §1: PENDING_REVIEW is a proposal, not a load."""
    review = ingest(engineer, cleanup, "t04-pending.xlsx", make_workbook(unique_rows(2)))
    load_id = review["batch"]["load_id"]
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM wavescan WHERE load_id = %s", (load_id,))
        assert cur.fetchone()[0] == 0


def test_commit_inserts_the_measurements_and_they_become_visible(engineer, cleanup):
    review = ingest(engineer, cleanup, "t04-commit.xlsx", make_workbook(unique_rows(3)))
    load_id = review["batch"]["load_id"]
    assert review["committable"] is True

    response = engineer.post(f"/api/v1/batches/{load_id}/commit")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "COMMITTED"
    assert body["measurements_inserted"] == 3
    assert body["counts"]["balanced"] is True

    # T76 CHANGE: narrowed to this batch (`load_id`, T19) -- see above.
    listed = engineer.get(
        "/api/v1/measurements",
        params={"limit": 500, "plant": "AAT", "load_id": load_id},
    ).json()
    mine = [i for i in listed["items"] if i["load_id"] == load_id]
    assert len(mine) == 3, "committed rows must appear in the T03 list endpoint"


def test_the_balance_check_holds_on_a_committed_batch(engineer, cleanup, admin_conn):
    """`02-PIPELINE.md` §6. The database CHECK is the guard, so this asserts
    the row satisfies it after commit rather than trusting the service."""
    review = ingest(engineer, cleanup, "t04-balance.xlsx", make_workbook(unique_rows(4)))
    load_id = review["batch"]["load_id"]
    assert engineer.post(f"/api/v1/batches/{load_id}/commit").status_code == 200
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT rows_in_file = rows_loaded + rows_duplicate + rows_rejected "
            "     + rows_quarantine "
            "FROM load_batch WHERE load_id = %s",
            (load_id,),
        )
        assert cur.fetchone()[0] is True


def test_commit_writes_an_audit_event_with_an_actor(engineer, cleanup, admin_conn):
    """`AGENTS.md` §2: the audit row is in the same transaction as the
    mutation, and since revision 0003 it names a person."""
    review = ingest(engineer, cleanup, "t04-audit.xlsx", make_workbook(unique_rows(2)))
    load_id = review["batch"]["load_id"]
    assert engineer.post(f"/api/v1/batches/{load_id}/commit").status_code == 200
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT action, actor_user_id, actor_type, after "
            "FROM audit_event WHERE entity_type='load_batch' AND entity_id=%s "
            "ORDER BY id",
            (str(load_id),),
        )
        events = cur.fetchall()
    actions = [e[0] for e in events]
    assert actions == ["batch.upload", "batch.commit"]
    for action, actor, actor_type, _ in events:
        assert actor is not None, f"{action} has no actor"
        assert actor_type == "user"
    assert events[-1][3]["measurements_inserted"] == 2


def test_the_uploader_is_recorded_on_the_batch(engineer, cleanup, admin_conn):
    review = ingest(engineer, cleanup, "t04-loadedby.xlsx", make_workbook(unique_rows(1)))
    who = review["batch"]["loaded_by"]
    assert who, "load_batch.loaded_by must be an app_user id, not a database role"
    with admin_conn.cursor() as cur:
        cur.execute("SELECT upn FROM app_user WHERE id = %s", (who,))
        assert cur.fetchone() is not None


def test_committing_twice_is_refused(engineer, cleanup):
    review = ingest(engineer, cleanup, "t04-twice.xlsx", make_workbook(unique_rows(2)))
    load_id = review["batch"]["load_id"]
    assert engineer.post(f"/api/v1/batches/{load_id}/commit").status_code == 200
    second = engineer.post(f"/api/v1/batches/{load_id}/commit")
    assert second.status_code == 409
    assert "already committed" in second.json()["message"]


def test_a_missing_batch_is_a_404(engineer):
    assert engineer.post("/api/v1/batches/99999999/commit").status_code == 404


# =====================================================================
# 2. identity and RBAC
# =====================================================================
def test_the_first_authenticated_request_creates_an_app_user(engineer, admin_conn):
    """T03 shipped `user_id: null` on purpose. T04 is where it stops being
    null, because revision 0003 grants the INSERT that makes it possible."""
    body = engineer.get("/api/v1/auth/me").json()
    assert body["principal"]["user_id"], "principal.user_id is still null"
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT upn, display_name FROM app_user WHERE id = %s",
            (body["principal"]["user_id"],),
        )
        assert cur.fetchone() is not None


def test_a_viewer_cannot_upload(viewer):
    response = upload(viewer, "t04-viewer.xlsx", make_workbook(unique_rows(1)))
    assert response.status_code == 403
    assert "engineer" in response.json()["message"]


def test_a_viewer_cannot_commit(viewer):
    assert viewer.post("/api/v1/batches/1/commit").status_code == 403


def test_a_viewer_creates_no_batch_when_refused(viewer, admin_conn):
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        before = cur.fetchone()[0]
    upload(viewer, "t04-viewer-2.xlsx", make_workbook(unique_rows(1)))
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        assert cur.fetchone()[0] == before


# =====================================================================
# 3. door checks -- refusal leaves no trace
# =====================================================================
def test_an_oversized_file_is_capped_before_it_is_parsed(ingest_apps, admin_conn):
    """The route must not `read()` the whole part and only then refuse.

    A 2 GB file named .xlsx would otherwise become a 2 GB `bytes` object
    in the API process; `02-PIPELINE.md` §2 wants the door shut first.
    The cap is the configured limit + 1 so check_file can still name
    TOO_LARGE."""
    client = ingest_apps(upload_max_file_bytes=64, ingest_worker_enabled=False)
    data = make_workbook(unique_rows(1))
    assert len(data) > 64
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        before = cur.fetchone()[0]
    response = upload(client, "huge.xlsx", data)
    assert response.status_code == 400, response.text
    assert "TOO_LARGE" in response.json()["message"]
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        assert cur.fetchone()[0] == before


@pytest.mark.parametrize(
    "name,data,expected_reason",
    [
        ("script.xlsx", b"#!/bin/sh\n", "NOT_A_WORKBOOK"),
        ("book.xls", b"PK\x03\x04rest", "BAD_EXTENSION"),
        ("empty.xlsx", b"", "EMPTY_FILE"),
    ],
)
def test_a_hostile_file_is_refused_with_400(engineer, admin_conn, name, data, expected_reason):
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        before = cur.fetchone()[0]

    response = upload(engineer, name, data)
    assert response.status_code == 400, response.text
    assert expected_reason in response.json()["message"]

    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        assert cur.fetchone()[0] == before, (
            "a refused upload must not create a load_batch "
            "(docs/02-PIPELINE.md section 2)"
        )


def test_the_error_shape_survives_a_door_check(engineer):
    body = upload(engineer, "script.xlsx", b"nope").json()
    assert set(body) == {"statusCode", "error", "message", "correlationId"}


# =====================================================================
# 4. the reload guard
# =====================================================================
def test_identical_bytes_are_refused_with_409_naming_the_prior_batch(engineer, cleanup):
    data = make_workbook(unique_rows(2))
    first = upload(engineer, "t04-dup.xlsx", data)
    assert first.status_code == 202
    load_id = first.json()["batches"][0]["load_id"]
    cleanup.append(load_id)
    wait_for(engineer, load_id)

    again = upload(engineer, "t04-dup-renamed.xlsx", data)
    assert again.status_code == 409
    assert f"load_id {load_id}" in again.json()["message"], (
        "the refusal must name the batch that already holds these bytes"
    )


def test_the_reload_guard_is_a_database_constraint(app_conn, admin_conn):
    """`02-PIPELINE.md` §3: "a database constraint, not only an application
    SELECT". Proven by going around the application entirely."""
    digest = uuid.uuid4().hex + uuid.uuid4().hex
    with app_conn.cursor() as cur:
        cur.execute(
            "INSERT INTO load_batch (source_file, source_sheet, file_hash, "
            "raw_store_key, status) VALUES ('a.xlsx','Sheet1',%s,'k','RUNNING')",
            (digest,),
        )
        with pytest.raises(psycopg.errors.UniqueViolation):
            cur.execute(
                "INSERT INTO load_batch (source_file, source_sheet, file_hash, "
                "raw_store_key, status) "
                "VALUES ('b.xlsx','Sheet1',%s,'k','RUNNING')",
                (digest,),
            )


def test_two_sheets_of_one_file_are_two_batches(engineer, cleanup):
    """The reason the unique key is (file_hash, source_sheet) and not
    file_hash: a measurement is identified by (load_id, source_row), so two
    sheets that both have a row 2 cannot share a batch. Revision 0003."""
    data = make_workbook([], sheets={
        "Sheet A": unique_rows(2),
        "Sheet B": unique_rows(2),
    })
    response = upload(engineer, "t04-two-sheets.xlsx", data)
    assert response.status_code == 202, response.text
    sheets = response.json()["batches"]
    assert {s["sheet"] for s in sheets} == {"Sheet A", "Sheet B"}
    for entry in sheets:
        cleanup.append(entry["load_id"])
        assert wait_for(engineer, entry["load_id"])["batch"]["status"] == "PENDING_REVIEW"


# =====================================================================
# 5. templates and quarantine
# =====================================================================
def test_an_unknown_template_is_rejected_without_parsing_the_rows(engineer, cleanup, admin_conn):
    """`02-PIPELINE.md` §4."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Style", "Orientation", "Panel", "LW", "SW", "CF"])
    ws.append(["P703", "H", "Hood", 1.0, 2.0, 60.0])
    buffer = io.BytesIO()
    wb.save(buffer)

    response = upload(engineer, "t04-unknown-template.xlsx", buffer.getvalue())
    assert response.status_code == 202
    entry = response.json()["batches"][0]
    load_id = entry["load_id"]
    cleanup.append(load_id)
    assert entry["status"] == "REJECTED_TEMPLATE"

    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM stg_raw WHERE load_id = %s", (load_id,))
        assert cur.fetchone()[0] == 0, "no data row may be parsed"
        cur.execute("SELECT count(*) FROM wavescan WHERE load_id = %s", (load_id,))
        assert cur.fetchone()[0] == 0
        cur.execute(
            "SELECT kind, status, payload FROM proposed_mapping WHERE load_id = %s",
            (load_id,),
        )
        row = cur.fetchone()
    assert row is not None, "an Admin needs a proposal to act on"
    assert row[0] == "COLUMN" and row[1] == "PROPOSED"
    assert "plant" in row[2]["missing_required"]


def test_a_proposed_mapping_is_never_applied_automatically(engineer, cleanup, admin_conn):
    """`02-PIPELINE.md` §4: "never auto-apply".

    T08 CHANGE: this used to read "no proposal anywhere is anything but
    PROPOSED", which was true while nothing could approve one. T08 adds an
    Admin route that can, so the old form would now fail for the right
    reason -- and asserting it globally would also make this test depend on
    what other suites had run.

    The invariant it was protecting is unchanged and is what is asserted
    instead: **ingest** never decides a mapping. Every proposal ingest itself
    writes is PROPOSED, applied by nothing, with no alias to show for it.
    Approval is a separate route, a different role, and T08's own suite.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Mystery Column", "CF"])
    ws.append(["anything", 60.0])
    buffer = io.BytesIO()
    wb.save(buffer)

    response = upload(engineer, "t04-never-applied.xlsx", buffer.getvalue())
    load_id = response.json()["batches"][0]["load_id"]
    cleanup.append(load_id)

    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT status, decided_by, decided_at FROM proposed_mapping "
            "WHERE load_id = %s",
            (load_id,),
        )
        rows = cur.fetchall()
        cur.execute(
            "SELECT count(*) FROM ref_column_alias WHERE raw_header = 'mystery column'"
        )
        aliases = cur.fetchone()[0]

    assert rows, "ingest wrote no proposal to check"
    assert all(r == ("PROPOSED", None, None) for r in rows), rows
    assert aliases == 0, "ingest wrote an alias"


def test_an_unknown_value_becomes_a_reject_not_a_guess(engineer, cleanup):
    rows = unique_rows(2)
    rows[0]["Color"] = "Fictional Purple"
    review = ingest(engineer, cleanup, "t04-reject.xlsx", make_workbook(rows))
    assert review["batch"]["counts"]["rows_rejected"] == 1
    assert review["reject_reasons"] == {"UNKNOWN_COLOR": 1}
    assert review["batch"]["counts"]["balanced"] is True


def test_colour_not_offered_at_this_plant_is_held(engineer, cleanup):
    """Traction Green is offered at AAT, not FTM. Identity still resolves."""
    data = make_workbook(unique_rows(1, Plant="FTM", Color="3HC Traction Green"))
    response = upload(engineer, "t04-not-offered.xlsx", data, "FTM")
    assert response.status_code == 202, response.text
    load_id = response.json()["batches"][0]["load_id"]
    cleanup.append(load_id)
    review = wait_for(engineer, load_id)
    assert review["reject_reasons"].get("COLOR_NOT_OFFERED") == 1, review.get("reject_reasons")
    held = [h for h in review["held_labels"] if h["reason"] == "COLOR_NOT_OFFERED"]
    assert held, review["held_labels"]
    assert held[0]["display_group"] == "plant_offer"
    # T39 CHANGE (ADR 0013 decision 10): an Engineer matches the spelling
    # onto an existing catalogue name; confirming one FTM does not list yet
    # offers it there. Asked every time, official names, no default.
    assert held[0]["teachable_kind"] == "ALIAS_COLOR"
    assert "Traction Green" in held[0]["candidates"]
    assert held[0]["candidates"] == sorted(held[0]["candidates"]), "no default first"
    assert "will not load yet" in held[0]["operator_message"]


def test_an_unnamed_model_is_quarantined(engineer, cleanup):
    rows = unique_rows(2)
    rows[0]["Model"] = "(none)"
    review = ingest(engineer, cleanup, "t04-quarantine.xlsx", make_workbook(rows))
    assert review["batch"]["counts"]["rows_quarantine"] == 1
    assert review["unidentified_rows"] == 1


def test_a_committed_batch_never_contains_a_rejected_row(
    engineer, admin, cleanup, admin_conn
):
    """The claim is unchanged: a rejected row does not become a measurement.

    T19 CHANGE: the commit now needs the rejected rows acknowledged first,
    so the test acknowledges them the way an Admin would. The refusal before
    the acknowledgement is asserted too -- dropping that would leave the gate
    untested here and make this read as if the gate did not exist.
    """
    rows = unique_rows(3)
    rows[0]["Color"] = "Fictional Purple"
    review = ingest(engineer, cleanup, "t04-partial.xlsx", make_workbook(rows))
    load_id = review["batch"]["load_id"]
    assert review["batch"]["counts"]["rows_rejected"] == 1
    assert review["committable"] is False
    assert "WARNINGS_UNACKNOWLEDGED" in review["commit_blockers"]

    refused = engineer.post(f"/api/v1/batches/{load_id}/commit")
    assert refused.status_code == 409, refused.text
    assert "WARNINGS_UNACKNOWLEDGED" in refused.json()["message"]

    assert acknowledge(admin, review, "REJECTED_ROWS").status_code == 201
    assert engineer.post(f"/api/v1/batches/{load_id}/commit").status_code == 200
    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM wavescan WHERE load_id = %s", (load_id,))
        assert cur.fetchone()[0] == 2


# =====================================================================
# 6. D5 -- FVL
# =====================================================================
def test_a_file_from_an_un_onboarded_plant_is_refused(engineer, admin_conn):
    """D5 stays open: FVL is refused with one clear error, and no zone-14
    hack. Driven by ref_plant.is_active, so the test seeds the situation
    rather than naming a plant the code knows about."""
    with admin_conn.cursor() as cur:
        cur.execute("SELECT plant_code FROM ref_plant WHERE NOT is_active LIMIT 1")
        row = cur.fetchone()
    if row is None:
        pytest.skip("no inactive plant in ref_plant to exercise the gate")
    inactive = row[0]

    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        before = cur.fetchone()[0]

    data = make_workbook([{**r, "Plant": inactive} for r in unique_rows(2)])
    # T17: the upload declares the same inactive plant its cells name, so the
    # only thing left to refuse is the onboarding decision itself. (A workbook
    # whose cells disagreed with the claim would get the more precise
    # wrong-plant 400 first -- tests/test_t17_expected_plant.py.)
    response = upload(engineer, "t04-inactive-plant.xlsx", data, expected_plant=inactive)
    assert response.status_code == 422, response.text
    message = response.json()["message"]
    assert inactive in message and "D5" in message

    with admin_conn.cursor() as cur:
        cur.execute("SELECT count(*) FROM load_batch")
        assert cur.fetchone()[0] == before, "a refused plant must leave no batch"


# =====================================================================
# 7. the worker
# =====================================================================
@pytest.fixture()
def paused_worker(engineer, monkeypatch):
    """Stop the background worker for the duration of one test.

    Two tests below drive the claim by hand. With a worker running they
    would race it for the same job and fail intermittently, which is a
    test-design problem rather than a fact about the queue.

    T76 CHANGE: stopping *this* module's worker was not enough. Every test
    module that imports ``engineer`` builds its own app with its own worker
    thread, and each claims **any** QUEUED job -- so another module's worker
    took this test's job. ``run_once`` calls ``ingest_jobs.claim`` through the
    module, so patching it pauses every in-process worker until the test
    ends; the test's own claim is plain SQL and is not affected. A worker in
    another process (standing 8012) cannot be paused from here -- stop it
    for test runs.
    """
    worker = engineer.app.state.ingest_worker
    worker.stop()
    monkeypatch.setattr(ingest_jobs, "claim", lambda conn: None)
    yield worker
    worker.start()


def test_two_workers_never_claim_the_same_job(app_conn, admin_conn, cleanup, paused_worker):
    """`02-PIPELINE.md` §9. SKIP LOCKED is what makes N workers safe: a job
    another transaction holds is passed over, not waited on."""
    with admin_conn.cursor() as cur:
        cur.execute(
            "INSERT INTO load_batch (source_file, source_sheet, file_hash, "
            "raw_store_key, status) VALUES ('claim.xlsx','S',%s,'k','RUNNING') "
            "RETURNING load_id",
            (uuid.uuid4().hex + uuid.uuid4().hex,),
        )
        load_id = cur.fetchone()[0]
        cleanup.append(load_id)
        cur.execute(
            "INSERT INTO ingest_job (load_id, status) VALUES (%s,'QUEUED')", (load_id,)
        )

    dsn = _dsn_of(app_conn)
    claim = (
        "SELECT job_id FROM ingest_job WHERE status='QUEUED' AND load_id = %s "
        "ORDER BY job_id FOR UPDATE SKIP LOCKED LIMIT 1"
    )
    with psycopg.connect(dsn) as first, psycopg.connect(dsn) as second:
        with first.cursor() as a, second.cursor() as b:
            a.execute(claim, (load_id,))
            taken = a.fetchone()
            b.execute(claim, (load_id,))
            missed = b.fetchone()
        first.rollback()
        second.rollback()

    assert taken is not None, "the first worker did not claim the queued job"
    assert missed is None, "the second worker took a job the first was holding"


def _dsn_of(conn) -> str:
    info = conn.info
    return f"host={info.host} port={info.port} dbname={info.dbname} user={info.user}"


def test_a_job_that_cannot_be_processed_fails_the_batch_rather_than_hanging(
    admin_conn, engineer, cleanup
):
    """A batch whose bytes are missing must end in a status a reviewer can
    see, not sit in RUNNING for ever."""
    with admin_conn.cursor() as cur:
        cur.execute(
            "INSERT INTO load_batch (source_file, source_sheet, file_hash, "
            "raw_store_key, status) VALUES ('gone.xlsx','S',%s,%s,'RUNNING') "
            "RETURNING load_id",
            (uuid.uuid4().hex + uuid.uuid4().hex, "wavescan/raw/" + "0" * 64),
        )
        load_id = cur.fetchone()[0]
        cleanup.append(load_id)
        cur.execute(
            "INSERT INTO ingest_job (load_id, status) VALUES (%s,'QUEUED')", (load_id,)
        )

    body = wait_for(engineer, load_id, timeout=40)
    assert body["batch"]["status"] == "FAILED"
    assert "missing" in (body["batch"]["notes"] or "").lower()


def test_the_worker_helpers_are_importable_as_a_standalone_process():
    """`python -m app.workers.ingest_worker` is the deployment story for a
    machine where the API and the worker are separate processes."""
    assert callable(ingest_worker.main)
    assert callable(ingest_worker.run_once)
    assert ingest_jobs.MAX_ATTEMPTS >= 1


# =====================================================================
# 8. T01's guarantees still hold after ingest exists
# =====================================================================
def test_the_ingest_role_still_cannot_delete_or_rewrite_a_measurement(app_conn):
    """The whole point of doing T01 first. Ingest inserts; it never needed
    either privilege, and it still does not have them."""
    with app_conn.cursor() as cur:
        for statement in (
            "DELETE FROM wavescan",
            "UPDATE wavescan SET cf = 1",
            "DELETE FROM load_batch",
            "UPDATE wavescan SET is_superseded = true",
        ):
            with app_conn.transaction():
                with pytest.raises(psycopg.errors.InsufficientPrivilege):
                    cur.execute(statement)


def test_an_actor_is_mandatory_for_a_supersession(app_conn, sample_measurement_id):
    """Revision 0003, the constraint T02's review asked T03 to defer here."""
    with app_conn.cursor() as cur:
        with pytest.raises(psycopg.errors.NotNullViolation):
            cur.execute(
                "INSERT INTO measurement_supersession "
                "(superseded_measurement_id, reason, actor_user_id) "
                "VALUES (%s, 'no actor', NULL)",
                (sample_measurement_id,),
            )


def test_ingest_never_inserts_a_row_that_is_already_retired(admin_conn):
    """A T02 trigger refuses it, for the owner too. Asserted here because
    T04 is the first code that inserts measurements at all."""
    with admin_conn.cursor() as cur:
        # Any retired row must have got that way through a supersession record.
        cur.execute(
            """
            SELECT count(*) FROM wavescan w
             WHERE w.is_superseded
               AND NOT EXISTS (SELECT 1 FROM measurement_supersession m
                                WHERE m.superseded_measurement_id = w.measurement_id)
            """
        )
        assert cur.fetchone()[0] == 0


def test_a_committed_batch_is_visible_but_a_pending_one_is_not(engineer, cleanup):
    """The active-data predicate, exercised through ingest rather than
    through a hand-made fixture row."""
    pending = ingest(engineer, cleanup, "t04-hidden.xlsx", make_workbook(unique_rows(2)))
    committed = ingest(engineer, cleanup, "t04-shown.xlsx", make_workbook(unique_rows(2)))
    committed_id = committed["batch"]["load_id"]
    assert engineer.post(f"/api/v1/batches/{committed_id}/commit").status_code == 200

    # T76 CHANGE: ask for each batch by `load_id` (T19 filter) instead of
    # scanning the first 500 of every live measurement -- the standing
    # database holds thousands, so the newest batch need not be on that page.
    shown = engineer.get(
        "/api/v1/measurements", params={"limit": 500, "load_id": committed_id}
    ).json()
    assert shown["page"]["total"] == 2
    assert {i["load_id"] for i in shown["items"]} == {committed_id}
    hidden = engineer.get(
        "/api/v1/measurements",
        params={"limit": 500, "load_id": pending["batch"]["load_id"]},
    ).json()
    assert hidden["page"]["total"] == 0 and hidden["items"] == []


# =====================================================================
# 9. a real plant file
# =====================================================================
AAT_JUNE = VAULT / "AAT" / "06 Jun'26 CF data AAT.xlsx"


@pytest.mark.skipif(not AAT_JUNE.is_file(), reason="reference vault not reachable")
def test_a_real_aat_workbook_ingests_and_commits(
    engineer, admin, cleanup, admin_conn
):
    """The evidence the T04 brief asks for: one real AAT file, uploaded and
    committed through the API, never copied into this repository.

    The counts are asserted as relationships, not as literals: this database
    may already hold some of these vehicles, which is exactly when the
    cross-file duplicate rule fires.
    """
    data = AAT_JUNE.read_bytes()
    forget_bytes(admin_conn, data)
    review = ingest(engineer, cleanup, AAT_JUNE.name, data)
    batch = review["batch"]
    assert batch["status"] == "PENDING_REVIEW"
    counts = batch["counts"]
    assert counts["rows_in_file"] == 629, "non-blank data rows in the June file"
    assert counts["balanced"] is True
    assert counts["rows_loaded"] > 0

    # T19: a real file rejects or duplicates some rows, and those have to be
    # read and accepted before it commits. Which categories apply depends on
    # what this database already holds, so they are taken from the counts.
    acknowledge_required(admin, review)

    response = engineer.post(f"/api/v1/batches/{batch['load_id']}/commit")
    assert response.status_code == 200, response.text
    assert response.json()["measurements_inserted"] == counts["rows_loaded"]

    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT count(*), count(DISTINCT serial_no), min(cf), max(cf) "
            "FROM wavescan WHERE load_id = %s",
            (batch["load_id"],),
        )
        stored, vehicles, low, high = cur.fetchone()
    assert stored == counts["rows_loaded"]
    assert vehicles > 1
    assert 0 <= low <= high <= 100

    # The file stores 2.7 as 2.70000004768372 in 273 of its 3942 measurement
    # cells -- a float4 artefact of the source system. The vault loader
    # rounded those away on ingest; `AGENTS.md` §4 forbids that, so they are
    # stored as the file wrote them and rounded for display instead.
    with admin_conn.cursor() as cur:
        cur.execute(
            "SELECT count(*) FROM wavescan WHERE load_id = %s AND scale(lw) > 3",
            (batch["load_id"],),
        )
        assert cur.fetchone()[0] > 0, (
            "every LW came back with three decimals or fewer, so something "
            "rounded the file's own digits on the way in"
        )
